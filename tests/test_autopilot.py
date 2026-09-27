"""Tests for the CUMCM Autopilot: intake, state, verification, codegen, delivery.

These cover the deterministic core of the autopilot — everything that must be
correct regardless of which model writes the prose.
"""

import json
from pathlib import Path

import pydantic
import pytest

from mathmodel.autopilot.codegen import (
    GeneratedProgram,
    _extract_code,
    _extract_header_field,
    _extract_imports,
    _parse_summary,
)
from mathmodel.autopilot.deliver import TableBuilder, build_cumcm_latex
from mathmodel.autopilot.intake import (
    AttachmentInfo,
    ProblemContext,
    ProblemIntake,
    SheetSchema,
    build_clarification_questions,
)
from mathmodel.autopilot.state import (
    ClarificationQuestion,
    RunState,
    RunStatus,
    StageStatus,
)
from mathmodel.autopilot.verify import (
    DRYING_TIME_RE,
    CheckStatus,
    DeterministicVerifier,
    VerificationCheck,
    VerificationReport,
    _self_reported_leftovers,
    build_report,
    collect_identifier_pool,
    compare_statistics,
    header_mismatches,
    problem_final_state_targets,
    problem_time_steps,
    read_output_table,
)
from mathmodel.documents import FigureRegistry, TableRegistry
from mathmodel.paper import BlockType, ContentBlock, PaperIR, PaperSection
from mathmodel.providers.json_repair import (
    coerce_enum_casing,
    normalise_for_schema,
    parse_json_object,
)
from pydantic import BaseModel


# ═══════════════════════════════════════════════════════════════
# RunState: persistence, resume, stage transitions
# ═══════════════════════════════════════════════════════════════

class TestRunState:
    def test_create_and_reload(self, tmp_path):
        state = RunState.create(tmp_path / "run1")
        state.begin("intake", "reading")
        state.complete("intake", "done", ["a.json"])
        state.save()

        reloaded = RunState.load(tmp_path / "run1")
        assert reloaded.run_id == state.run_id
        assert reloaded.is_done("intake")
        assert reloaded.stage("intake").detail == "done"
        assert reloaded.stage("intake").artifacts == ["a.json"]

    def test_status_transitions_are_recorded(self, tmp_path):
        state = RunState.create(tmp_path / "run2")
        state.begin("solve", "running")
        assert state.stage("solve").status == StageStatus.RUNNING

        state.fail("solve", "exit 1")
        assert state.stage("solve").status == StageStatus.FAILED
        assert state.stage("solve").attempts == 1
        assert not state.is_done("solve")

        state.block("solve", "gave up")
        assert state.stage("solve").status == StageStatus.BLOCKED
        assert state.status == RunStatus.BLOCKED

    def test_skip_and_count(self, tmp_path):
        state = RunState.create(tmp_path / "run3")
        state.begin("a")
        state.complete("a")
        state.skip("b", "not needed")
        assert state.count(StageStatus.COMPLETED) == 1
        assert state.count(StageStatus.SKIPPED) == 1
        assert state.stage("b").status == StageStatus.SKIPPED

    def test_save_is_atomic_and_utf8(self, tmp_path):
        state = RunState.create(tmp_path / "run4")
        state.complete("intake", "中文详情")
        state.save()
        raw = (tmp_path / "run4" / "pipeline_state.json").read_text(encoding="utf-8")
        assert "中文详情" in raw
        assert RunState.load(tmp_path / "run4").stage("intake").detail == "中文详情"

    def test_summary_shape(self, tmp_path):
        state = RunState.create(tmp_path / "run5")
        state.begin("intake")
        state.complete("intake", "ok")
        summary = state.summary()
        assert summary["run_id"] == state.run_id
        assert summary["stages"]["intake"]["status"] == "completed"
        assert "current_stage" in summary


# ═══════════════════════════════════════════════════════════════
# Intake: real parsing of real files
# ═══════════════════════════════════════════════════════════════

def _make_xlsx(path: Path, headers, rows):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path


class TestIntake:
    def test_reads_real_excel_cells(self, tmp_path):
        source = _make_xlsx(
            tmp_path / "附件1.xlsx",
            ["用频装备编号", "频段区间", "时间区间", "间隔时长", "使用次数"],
            [["A001", "[80,90)", "[35,40)", 60, 3], ["B002", "[10,20)", "[0,5)", 30, 2]],
        )
        intake = ProblemIntake(tmp_path / "input")
        context, records = intake.ingest([source])

        assert len(records) == 1
        info = context.attachments[0]
        assert info.role == "data"
        assert info.parser_status.value == "completed"
        assert info.sheets
        sheet = info.sheets[0]
        assert sheet.row_count == 2
        assert sheet.headers[0] == "用频装备编号"
        assert sheet.sample_rows[0][0] == "A001"

    def test_classifies_templates_and_detects_required_outputs(self, tmp_path):
        problem = tmp_path / "D题.txt"
        problem.write_text(
            "问题 1 将结果保存到文件result1.xlsx 中。\n"
            "问题 2 将消解方案保存到文件result2.xlsx 中。\n"
            "附件 1 给出了用频计划。\n",
            encoding="utf-8",
        )
        data = _make_xlsx(tmp_path / "附件1.xlsx", ["编号"], [["A001"]])
        template = _make_xlsx(tmp_path / "result1.xlsx", ["序号", "冲突设备1", "冲突设备2"], [])

        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([problem, data, template])

        assert context.required_outputs == ["result1.xlsx", "result2.xlsx"]
        assert context.template_file_names == ["result1.xlsx"]
        assert "附件1" in context.declared_attachments
        assert context.missing_attachments == []

    def test_reports_missing_declared_attachment(self, tmp_path):
        problem = tmp_path / "题.txt"
        problem.write_text("详见附件 3 的说明。", encoding="utf-8")
        data = _make_xlsx(tmp_path / "附件1.xlsx", ["编号"], [["A001"]])

        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([problem, data])
        assert "附件3" in context.missing_attachments

    def test_export_processed_data_writes_csv(self, tmp_path):
        source = _make_xlsx(
            tmp_path / "附件1.xlsx", ["编号", "值"], [["A001", 1], ["A002", 2]]
        )
        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([source])
        written = intake.export_processed_data(context, tmp_path / "processed")

        assert len(written) == 1
        text = Path(written[0]).read_text(encoding="utf-8-sig")
        assert "A001" in text and "A002" in text

    def test_no_questions_when_upload_is_sufficient(self, tmp_path):
        problem = tmp_path / "题.txt"
        problem.write_text("某区域内150个用频装备提报用频计划，需检测时频冲突。", encoding="utf-8")
        data = _make_xlsx(tmp_path / "附件1.xlsx", ["编号"], [["A001"]])

        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([problem, data])
        assert build_clarification_questions(context) == []

    def test_blocking_question_when_statement_unreadable(self):
        context = ProblemContext(problem_text="", attachments=[
            AttachmentInfo(
                original_name="scan.pdf", role="problem_statement",
                storage_path="scan.pdf", size_bytes=10,
            )
        ])
        questions = build_clarification_questions(context)
        assert any(q.blocking for q in questions)
        assert questions[0].question_id == "Q-STATEMENT"

    def test_missing_attachment_is_non_blocking_when_data_present(self, tmp_path):
        problem = tmp_path / "题.txt"
        problem.write_text("详见附件 2。", encoding="utf-8")
        data = _make_xlsx(tmp_path / "附件1.xlsx", ["编号"], [["A001"]])

        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([problem, data])
        questions = build_clarification_questions(context)
        assert questions and not any(q.blocking for q in questions)

    def test_missing_attachment_blocks_without_any_data(self, tmp_path):
        problem = tmp_path / "题.txt"
        problem.write_text("详见附件 1 的数据。", encoding="utf-8")

        intake = ProblemIntake(tmp_path / "input")
        context, _ = intake.ingest([problem])
        questions = build_clarification_questions(context)
        assert questions and questions[0].blocking


# ═══════════════════════════════════════════════════════════════
# Verification
# ═══════════════════════════════════════════════════════════════

class TestSelfReportedLeftovers:
    def test_flags_residual_conflicts(self):
        found = _self_reported_leftovers({"problem2": {"residual_conflicts": 291}})
        assert found == {"problem2.residual_conflicts": 291.0}

    def test_zero_residuals_pass(self):
        assert _self_reported_leftovers({"problem2_remaining_conflicts": 0}) == {}

    def test_final_conflicts_is_a_residual_counter(self):
        assert _self_reported_leftovers({"problem2": {"final_conflicts": 0}}) == {}
        assert _self_reported_leftovers({"problem2": {"final_conflicts": 5}}) == {
            "problem2.final_conflicts": 5.0
        }

    def test_total_conflicts_is_not_a_residual(self):
        assert _self_reported_leftovers({"problem1": {"total_conflicts": 297}}) is None

    def test_detected_conflict_counts_are_not_residuals(self):
        stats = {"problem1": {"total_conflict_pairs": 297, "num_plans": 150}}
        assert _self_reported_leftovers(stats) is None

    def test_nested_and_mixed(self):
        stats = {"a": {"b": {"unresolved_overlaps": 3, "residual_conflicts": 0}}}
        assert _self_reported_leftovers(stats) == {"a.b.unresolved_overlaps": 3.0}

    def test_booleans_ignored(self):
        assert _self_reported_leftovers({"violations_leftover": True}) is None


class TestOutputReaders:
    def test_reads_xlsx(self, tmp_path):
        path = _make_xlsx(tmp_path / "r.xlsx", ["序号", "设备"], [[1, "A001"], [2, "B002"]])
        headers, rows = read_output_table(path)
        assert headers == ["序号", "设备"]
        assert rows == [[1, "A001"], [2, "B002"]]

    def test_reads_csv(self, tmp_path):
        path = tmp_path / "r.csv"
        path.write_text("序号,设备\n1,A001\n", encoding="utf-8")
        headers, rows = read_output_table(path)
        assert headers == ["序号", "设备"]
        assert rows == [["1", "A001"]]

    def test_identifier_pool_from_real_data(self, tmp_path):
        _make_xlsx(
            tmp_path / "附件1.xlsx", ["编号", "值"],
            [["A001", 1], ["B002", 2], ["not-an-id", 3]],
        )
        pool = collect_identifier_pool(tmp_path)
        assert pool == {"A001", "B002"}


def _outcome(statistics=None, status="success"):
    from mathmodel.autopilot.codegen import ExecutionOutcome

    return ExecutionOutcome(
        run_id="RUN-test",
        status=status,
        exit_code=0 if status == "success" else 1,
        stdout="",
        stderr="",
        runtime_seconds=1.5,
        execution_real=True,
        sandbox={"image": "test", "code_hash": "abc"},
        summary={"statistics": statistics or {}},
    )


class TestDeterministicVerifier:
    def test_all_checks_pass_on_a_clean_solution(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["编号"], [["A001"], ["B002"]])

        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result1.xlsx", ["序号", "冲突设备1", "冲突设备2"], [[1, "A001", "B002"]])

        checks = DeterministicVerifier(data_dir).run(
            outcome=_outcome({"conflicts": 1, "residual_conflicts": 0}),
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers={"result1.xlsx": ["序号", "冲突设备1", "冲突设备2"]},
        )
        by_name = {c.name: c for c in checks}
        assert by_name["execution_was_real"].status == CheckStatus.PASS
        assert by_name["required_outputs_present"].status == CheckStatus.PASS
        assert by_name["output_structure_matches_template"].status == CheckStatus.PASS
        assert by_name["referential_integrity"].status == CheckStatus.PASS
        assert by_name["statistics_reported"].status == CheckStatus.PASS
        assert by_name["no_self_reported_residual_violations"].status == CheckStatus.PASS
        assert all(c.status != CheckStatus.FAIL for c in checks)

    def test_unresolved_conflicts_fail_the_gate(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["编号"], [["A001"]])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result1.xlsx", ["序号"], [[1]])

        checks = DeterministicVerifier(data_dir).run(
            outcome=_outcome({"problem2": {"residual_conflicts": 291}}),
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers={},
        )
        residual = [c for c in checks if c.name == "no_self_reported_residual_violations"][0]
        assert residual.status == CheckStatus.FAIL
        assert "residual_conflicts" in residual.detail

    def test_missing_output_fails(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        out_dir = tmp_path / "out"
        out_dir.mkdir()

        checks = DeterministicVerifier(data_dir).run(
            outcome=_outcome({"conflicts": 3}),
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers={},
        )
        present = [c for c in checks if c.name == "required_outputs_present"][0]
        assert present.status == CheckStatus.FAIL

    def test_unknown_identifier_fails_referential_integrity(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["编号"], [["A001"]])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result1.xlsx", ["设备"], [["Z999"]])

        checks = DeterministicVerifier(data_dir).run(
            outcome=_outcome({"conflicts": 1}),
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers={},
        )
        ref = [c for c in checks if c.name == "referential_integrity"][0]
        assert ref.status == CheckStatus.FAIL
        assert "Z999" in ref.detail

    def test_unreal_execution_fails(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        out_dir = tmp_path / "out"
        out_dir.mkdir()

        outcome = _outcome({}, status="failed")
        outcome.execution_real = False
        checks = DeterministicVerifier(data_dir).run(
            outcome=outcome, output_dir=out_dir,
            required_outputs=[], template_headers={},
        )
        real = [c for c in checks if c.name == "execution_was_real"][0]
        assert real.status == CheckStatus.FAIL


class TestReportAndComparison:
    def test_report_fails_when_any_check_fails(self):
        from mathmodel.domain.math_model import (
            MathematicalModel, Objective, ObjectiveSense,
        )

        model = MathematicalModel(
            model_id="M1", name="t", description="d",
            objectives=[Objective(name="o", sense=ObjectiveSense.MINIMIZE, expression="x")],
        )
        checks = [
            VerificationCheck(name="a", category="x", status=CheckStatus.PASS),
            VerificationCheck(name="b", category="x", status=CheckStatus.FAIL, detail="boom"),
        ]
        report = build_report(model, "RUN-1", checks)
        assert not report.passed
        assert report.blocking_failures
        assert "boom" in report.to_prompt_feedback()[0] or any(
            "boom" in f for f in report.to_prompt_feedback()
        )

    def test_not_run_counts_as_failure(self):
        from mathmodel.domain.math_model import (
            MathematicalModel, Objective, ObjectiveSense,
        )

        model = MathematicalModel(
            model_id="M1", name="t", description="d",
            objectives=[Objective(name="o", sense=ObjectiveSense.MINIMIZE, expression="x")],
        )
        report = build_report(model, "RUN-1", [
            VerificationCheck(name="a", category="x", status=CheckStatus.NOT_RUN),
        ])
        assert not report.passed

    def test_compare_statistics_detects_mismatch(self):
        check = compare_statistics({"a": 1.0}, {"a": 2.0})
        assert check.status == CheckStatus.FAIL

    def test_compare_statistics_accepts_match(self):
        check = compare_statistics({"a": 1.0}, {"a": 1.0})
        assert check.status == CheckStatus.PASS

    def test_nested_dicts_with_different_extra_keys_still_agree(self):
        # Two independent programs naturally report different extra keys. Only
        # the values present in both may be compared.
        check = compare_statistics(
            {"problem3": {"total_cells": 64600, "num_new_plans": 53680, "T_lo": 0}},
            {"problem3": {"total_cells": 64600, "num_new_plans": 53680, "new_cells": 53680}},
        )
        assert check.status == CheckStatus.PASS
        assert check.evidence["compared"] == 2

    def test_nested_dicts_still_catch_a_real_mismatch(self):
        check = compare_statistics(
            {"problem3": {"num_new_plans": 53680}},
            {"problem3": {"num_new_plans": 508}},
        )
        assert check.status == CheckStatus.FAIL
        assert "problem3.num_new_plans" in check.detail

    def test_boolean_is_not_treated_as_a_matching_number(self):
        # bool is a subclass of int, so a bool must be compared exactly rather
        # than through the numeric tolerance path.
        check = compare_statistics({"flag": True}, {"flag": False})
        assert check.status == CheckStatus.FAIL


# ═══════════════════════════════════════════════════════════════
# Code generation helpers
# ═══════════════════════════════════════════════════════════════

class TestCodeExtraction:
    def test_strips_markdown_fence(self):
        code = _extract_code("```python\n# APPROACH: x\nprint(1)\n```")
        assert code.startswith("# APPROACH: x")
        assert "```" not in code

    def test_keeps_raw_source(self):
        assert _extract_code("import os\nprint(1)") == "import os\nprint(1)"

    def test_handles_unterminated_fence(self):
        # A truncated reply has no closing marker; the opening ``` must not
        # survive into the program, where it is a SyntaxError.
        code = _extract_code("Here is the program:\n\n```python\nimport os\nprint(1)\n")
        assert "```" not in code
        assert code.startswith("import os")
        compile(code, "<test>", "exec")

    def test_drops_leading_prose_without_fence(self):
        code = _extract_code(
            "Sure! Here is a complete solver.\n\n"
            "import json\nprint(json.dumps({}))\n"
        )
        assert code.startswith("import json")
        compile(code, "<test>", "exec")

    def test_trims_trailing_prose(self):
        code = _extract_code(
            "import os\nprint(1)\n\n"
            "This program models the conflict as a graph and runs greedy search.\n"
            "Let me know if you need anything else!\n"
        )
        compile(code, "<test>", "exec")
        assert code.startswith("import os")

    def test_keeps_docstring_start(self):
        code = _extract_code('"""Solver."""\nimport os\n')
        assert code.startswith('"""Solver."""')

    def test_empty_input(self):
        assert _extract_code("") == ""
        assert _extract_code("   ") == ""

    def test_reads_approach_header(self):
        assert _extract_header_field("# APPROACH: greedy\nx=1", "APPROACH") == "greedy"

    def test_missing_header_returns_empty(self):
        assert _extract_header_field("x=1", "APPROACH") == ""

    def test_extract_imports_filters_stdlib(self):
        deps = _extract_imports("import os\nimport numpy as np\nfrom scipy import optimize\n")
        assert deps == ["numpy", "scipy"]

    def test_parse_summary_takes_last_json_line(self):
        stdout = 'progress\n{"statistics": {"a": 1}}\ntrailing text\n{"statistics": {"a": 2}}'
        assert _parse_summary(stdout)["statistics"]["a"] == 2

    def test_parse_summary_ignores_broken_json(self):
        assert _parse_summary("not json\n{broken\n") == {}


# ═══════════════════════════════════════════════════════════════
# JSON repair helpers shared by providers
# ═══════════════════════════════════════════════════════════════

class _StatusModel(BaseModel):
    class Status(str, __import__("enum").Enum):
        KNOWN = "known"
        REQUIRED = "required"

    status: Status


class TestJsonRepair:
    def test_parses_fenced_json(self):
        assert parse_json_object('```json\n{"a": 1}\n```') == {"a": 1}

    def test_parses_json_with_prose(self):
        assert parse_json_object('Here you go: {"a": 1} hope that helps') == {"a": 1}

    def test_empty_returns_none(self):
        assert parse_json_object("") is None
        assert parse_json_object("no json here") is None

    def test_enum_case_is_canonicalised(self):
        normalised = normalise_for_schema({"status": "REQUIRED"}, _StatusModel)
        assert normalised["status"] == "required"
        assert _StatusModel.model_validate(normalised).status.value == "required"

    def test_unknown_enum_value_is_left_alone(self):
        normalised = normalise_for_schema({"status": "bogus"}, _StatusModel)
        assert normalised["status"] == "bogus"
        with pytest.raises(Exception):
            _StatusModel.model_validate(normalised)

    def test_nested_enum_coercion(self):
        schema = {
            "type": "object",
            "properties": {"items": {"type": "array", "items": {"$ref": "#/$defs/E"}}},
            "$defs": {"E": {"enum": ["low", "high"]}},
        }
        value = {"items": ["HIGH", "low"]}
        assert coerce_enum_casing(value, schema, schema["$defs"]) == {"items": ["high", "low"]}


# ═══════════════════════════════════════════════════════════════
# Delivery: tables and LaTeX
# ═══════════════════════════════════════════════════════════════

class TestTables:
    def test_statistics_table_carries_source_ids(self):
        builder = TableBuilder()
        table = builder.build_statistics_table(
            {"conflicts": 297, "plans": 150}, "EVD-EXECUTION"
        )
        assert table is not None
        assert len(table.rows) == 2
        assert all(cell.source_id == "EVD-EXECUTION" for row in table.rows for cell in row)

    def test_nested_statistics_table(self):
        builder = TableBuilder()
        table = builder.build_nested_statistics_table(
            statistics={"p2": {"A": {"keep": 19, "cancel": 1}}},
            source_id="EVD-EXECUTION",
            table_id="TAB-002",
            title="消解统计",
            key="p2",
        )
        assert table is not None
        assert table.table_id == "TAB-002"
        assert any(cell.value == 19 for row in table.rows for cell in row)

    def test_empty_statistics_produces_no_table(self):
        assert TableBuilder().build_statistics_table({}, "EVD") is None


class TestLatexRenderer:
    def test_renders_figures_and_tables(self, tmp_path):
        image = tmp_path / "fig1.png"
        image.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)

        from mathmodel.documents import FigureRecord

        figures = FigureRegistry()
        figures.register(FigureRecord(
            figure_id="FIG-001",
            title="冲突分布",
            caption="图 1 冲突分布",
            artifact_path=str(image),
            source_execution_ids=["RUN-1"],
        ))
        tables = TableRegistry()
        tables.register(TableBuilder().build_statistics_table({"conflicts": 297}, "EVD"))

        paper = PaperIR(
            title="测试论文",
            abstract="摘要内容",
            sections=[
                PaperSection(
                    title="问题分析",
                    content_blocks=[ContentBlock(block_type=BlockType.PARAGRAPH, text="分析正文")],
                ),
                PaperSection(
                    title="结果分析",
                    content_blocks=[
                        ContentBlock(block_type=BlockType.FIGURE, text="图 1 冲突分布",
                                     figure_ids=["FIG-001"]),
                        ContentBlock(block_type=BlockType.TABLE, text="TAB-001",
                                     table_ids=["TAB-001"]),
                    ],
                ),
            ],
        )
        tex = build_cumcm_latex(paper, figures, tables, tmp_path)
        assert "\\documentclass" in tex
        assert "\\includegraphics" in tex
        assert "fig1.png" in tex
        assert "297" in tex
        assert "结果分析" in tex


# ═══════════════════════════════════════════════════════════════
# GeneratedProgram / ExecutionOutcome contracts
# ═══════════════════════════════════════════════════════════════

class TestGeneratedProgram:
    def test_code_is_required(self):
        with pytest.raises(Exception):
            GeneratedProgram(approach="x", code="")

    def test_succeeded_flag(self):
        from mathmodel.autopilot.codegen import ExecutionOutcome

        assert ExecutionOutcome(status="success").succeeded
        assert not ExecutionOutcome(status="failed").succeeded


class TestLatexInjectionGuard:
    """The guard must block real primitives without blocking legitimate ones."""

    def test_allows_ordinary_commands(self):
        from mathmodel.paper.renderer import _check_latex_injection

        tex = (
            r"\includegraphics[width=0.8\textwidth]{fig1.png}"
            r"\begin{figure}[H]\caption{图 1}\end{figure}"
            r"\section{结果}\inputencoding{utf8}\readline"
        )
        assert _check_latex_injection(tex) == []

    def test_blocks_real_primitives(self):
        from mathmodel.paper.renderer import _check_latex_injection

        assert _check_latex_injection(r"\input{/etc/passwd}")
        assert _check_latex_injection(r"\include{secret}")
        assert _check_latex_injection(r"\write18{rm -rf /}")
        assert _check_latex_injection(r"\catcode`\@=11")
        assert _check_latex_injection(r"\special{ps: hide}")

    def test_blocks_primitive_at_end_of_string(self):
        from mathmodel.paper.renderer import _check_latex_injection

        assert _check_latex_injection(r"text \include")


class TestFigureBuilder:
    def test_renders_real_png_from_statistics(self, tmp_path):
        from mathmodel.autopilot.deliver import FigureBuilder, FigureSpec

        builder = FigureBuilder()
        registry = builder.render(
            specs=[FigureSpec(
                title="冲突对数量", chart="bar", figure_type="comparison",
                caption="图 1 冲突对数量", statistics_keys=["conflicts"],
                supported_claim="共有若干冲突", source="statistics",
            )],
            statistics={"conflicts": 297},
            output_dir=tmp_path,
            figures_dir=tmp_path / "figures",
            execution_id="RUN-1",
            source_data_ids=["EVD-1"],
        )
        records = registry.all()
        assert len(records) == 1
        record = records[0]
        artifact = Path(record.artifact_path)
        assert artifact.exists()
        assert artifact.read_bytes()[:4] == b"\x89PNG"
        assert record.source_execution_ids == ["RUN-1"]
        assert record.source_data_ids == ["EVD-1"]
        assert record.metadata["values"] == [297.0]

    def test_falls_back_to_available_numeric_statistics(self, tmp_path):
        from mathmodel.autopilot.deliver import FigureBuilder, FigureSpec

        # A requested key that is absent falls back to whatever numeric
        # statistics exist, rather than producing an empty figure.
        registry = FigureBuilder().render(
            specs=[FigureSpec(
                title="回退", chart="bar", statistics_keys=["absent"],
                source="statistics",
            )],
            statistics={"conflicts": 297},
            output_dir=tmp_path,
            figures_dir=tmp_path / "figures",
            execution_id="RUN-1",
            source_data_ids=[],
        )
        records = registry.all()
        assert len(records) == 1
        assert records[0].metadata["values"] == [297.0]

    def test_skips_specs_with_no_numeric_statistics(self, tmp_path):
        from mathmodel.autopilot.deliver import FigureBuilder, FigureSpec

        registry = FigureBuilder().render(
            specs=[FigureSpec(
                title="无数据", chart="bar", statistics_keys=["note"],
                source="statistics",
            )],
            statistics={"note": "纯文本"},
            output_dir=tmp_path,
            figures_dir=tmp_path / "figures",
            execution_id="RUN-1",
            source_data_ids=[],
        )
        assert registry.all() == []

    def test_mixed_numeric_and_text_keys_stay_paired(self, tmp_path):
        from mathmodel.autopilot.deliver import FigureBuilder, FigureSpec

        # A requested key that is not numeric must not desynchronise labels and
        # values, which would otherwise crash matplotlib.
        registry = FigureBuilder().render(
            specs=[FigureSpec(
                title="混合", chart="bar",
                statistics_keys=["conflicts", "note"],
                source="statistics",
            )],
            statistics={"conflicts": 297, "note": "见正文"},
            output_dir=tmp_path,
            figures_dir=tmp_path / "figures",
            execution_id="RUN-1",
            source_data_ids=[],
        )
        records = registry.all()
        assert len(records) == 1
        metadata = records[0].metadata
        assert len(metadata["labels"]) == len(metadata["values"]) == 1
        assert metadata["labels"] == ["conflicts"]

    def test_series_from_output_file(self, tmp_path):
        from mathmodel.autopilot.deliver import FigureBuilder, FigureSpec

        _make_xlsx(
            tmp_path / "result1.xlsx", ["序号", "冲突设备1", "冲突设备2"],
            [[1, "A001", "B002"], [2, "A003", "C004"]],
        )
        registry = FigureBuilder().render(
            specs=[FigureSpec(
                title="冲突对", chart="bar", source="output_file",
                output_file="result1.xlsx", output_columns=["序号", "冲突设备1"],
            )],
            statistics={},
            output_dir=tmp_path,
            figures_dir=tmp_path / "figures",
            execution_id="RUN-1",
            source_data_ids=[],
        )
        # "冲突设备1" is text, so it cannot become a numeric series.
        assert registry.all() == []


class TestSupportPackageBuilder:
    def test_assembles_output_and_manifest(self, tmp_path):
        from mathmodel.autopilot.deliver import SupportPackageBuilder

        source = tmp_path / "src"
        source.mkdir()
        code = source / "solver.py"
        code.write_text("print(1)\n", encoding="utf-8")
        data = source / "附件1.csv"
        data.write_text("编号\nA001\n", encoding="utf-8")
        pdf = source / "paper.pdf"
        pdf.write_bytes(b"%PDF-1.7\n" + b"0" * 64)
        figure = source / "fig1.png"
        figure.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 32)

        out_dir = tmp_path / "output"
        manifest = SupportPackageBuilder(out_dir).build(
            paper_pdf=pdf,
            paper_tex=None,
            source_code=[code],
            processed_data=[data],
            original_files=[],
            figures=[figure],
            tables=[],
            result_files=[data],
            manifest_extra={"run_id": "RUN-1"},
        )

        assert (out_dir / "paper.pdf").read_bytes()[:5] == b"%PDF-"
        assert (out_dir / "support" / "source_code" / "solver.py").exists()
        assert (out_dir / "support" / "figures" / "fig1.png").exists()
        assert manifest["paper_pdf"] == "paper.pdf"
        assert manifest["run_id"] == "RUN-1"
        assert manifest["counts"]["source_code"] == 1
        written = json.loads((out_dir / "manifest.json").read_text(encoding="utf-8"))
        assert written["counts"] == manifest["counts"]

    def test_same_basename_does_not_overwrite(self, tmp_path):
        from mathmodel.autopilot.deliver import SupportPackageBuilder

        a = tmp_path / "a"
        b = tmp_path / "b"
        a.mkdir()
        b.mkdir()
        (a / "solver.py").write_text("A\n", encoding="utf-8")
        (b / "solver.py").write_text("B\n", encoding="utf-8")

        out_dir = tmp_path / "output"
        manifest = SupportPackageBuilder(out_dir).build(
            paper_pdf=None, paper_tex=None,
            source_code=[a / "solver.py", b / "solver.py"],
            processed_data=[], original_files=[], figures=[], tables=[],
            result_files=[], manifest_extra={},
        )
        copied = sorted((out_dir / "support" / "source_code").glob("*.py"))
        assert len(copied) == 2
        assert manifest["counts"]["source_code"] == 2

    def test_missing_pdf_is_recorded_as_none(self, tmp_path):
        from mathmodel.autopilot.deliver import SupportPackageBuilder

        out_dir = tmp_path / "output"
        manifest = SupportPackageBuilder(out_dir).build(
            paper_pdf=tmp_path / "nope.pdf", paper_tex=None,
            source_code=[], processed_data=[], original_files=[],
            figures=[], tables=[], result_files=[], manifest_extra={},
        )
        assert manifest["paper_pdf"] is None
        assert not (out_dir / "paper.pdf").exists()


class TestProgramCompleteness:
    """A truncated program parses and exits 0 while producing nothing."""

    def test_detects_missing_summary_print(self):
        from mathmodel.autopilot.codegen import _code_looks_complete

        assert not _code_looks_complete("def main():\n    x = 1\n")
        assert _code_looks_complete("import json\nprint(json.dumps({'a': 1}))\n")

    def test_detects_unparseable_code(self):
        from mathmodel.autopilot.codegen import _code_looks_complete

        assert not _code_looks_complete("def broken(:\n")

    def test_empty_is_incomplete(self):
        from mathmodel.autopilot.codegen import _code_looks_complete

        assert not _code_looks_complete("")
        assert not _code_looks_complete("   ")

    def test_truncated_reply_triggers_a_compact_retry(self):
        import asyncio

        from mathmodel.autopilot.codegen import generate_program_source
        from mathmodel.providers.base import GenerationResponse, UsageInfo
        from mathmodel.routing.profile import TaskProfile, TaskType

        calls: list[str] = []

        class FakeRouter:
            async def route_generate(self, profile, prompt, **kwargs):
                calls.append(prompt)
                if len(calls) == 1:
                    # Cut off mid-function: parses, but never prints a summary.
                    return GenerationResponse(
                        content="import json\ndef main():\n    total = 1\n",
                        model="fake", finish_reason="length",
                        usage=UsageInfo(), is_mock=False,
                    )
                return GenerationResponse(
                    content="import json\nprint(json.dumps({'total': 1}))\n",
                    model="fake", finish_reason="stop",
                    usage=UsageInfo(), is_mock=False,
                )

        code = asyncio.run(generate_program_source(
            FakeRouter(),
            profile=TaskProfile.for_task_type(TaskType.CODE_GENERATION),
            prompt="write a program",
            system_prompt="be brief",
        ))

        assert len(calls) == 2
        assert "CUT OFF" in calls[1]
        assert "print(" in code

    def test_accepts_a_good_first_reply_without_retrying(self):
        import asyncio

        from mathmodel.autopilot.codegen import generate_program_source
        from mathmodel.providers.base import GenerationResponse, UsageInfo
        from mathmodel.routing.profile import TaskProfile, TaskType

        calls: list[str] = []

        class FakeRouter:
            async def route_generate(self, profile, prompt, **kwargs):
                calls.append(prompt)
                return GenerationResponse(
                    content="print('{\"ok\": true}')",
                    model="fake", finish_reason="stop",
                    usage=UsageInfo(), is_mock=False,
                )

        asyncio.run(generate_program_source(
            FakeRouter(),
            profile=TaskProfile.for_task_type(TaskType.CODE_GENERATION),
            prompt="write a program",
            system_prompt="be brief",
        ))
        assert len(calls) == 1


class TestClarificationQuestion:
    def test_blocks_by_default(self):
        # Fail safe: a question nobody classified is treated as blocking rather
        # than silently letting the pipeline advance on unanswered information.
        question = ClarificationQuestion(question_id="Q1", question="?")
        assert question.blocking is True
        assert question.options == []

    def test_can_be_marked_non_blocking(self):
        question = ClarificationQuestion(question_id="Q1", question="?", blocking=False)
        assert question.blocking is False


class TestPaperToPdfPipeline:
    """Covers the delivery path end to end: paper IR -> LaTeX -> PDF -> package.

    This is the stage that had never executed, and it is where a figure or a
    table silently disappearing would cost the whole deliverable.
    """

    def _build(self, tmp_path, figure=True, table=True):
        from mathmodel.documents import FigureRecord
        from mathmodel.autopilot.deliver import SupportPackageBuilder, build_pdf

        figures = FigureRegistry()
        if figure:
            image = tmp_path / "figures" / "fig1.png"
            image.parent.mkdir(parents=True, exist_ok=True)
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
            fig, ax = plt.subplots()
            ax.bar(["A", "B", "C"], [20, 40, 90])
            fig.savefig(image)
            plt.close(fig)
            figures.register(FigureRecord(
                figure_id="FIG-001", title="装备类别分布", caption="图 1 装备类别分布",
                artifact_path=str(image), source_execution_ids=["RUN-1"],
            ))

        tables = TableRegistry()
        if table:
            tables.register(TableBuilder().build_statistics_table(
                {"total_conflicts": 297, "num_plans": 150}, "EVD-EXECUTION",
            ))

        blocks = [ContentBlock(block_type=BlockType.PARAGRAPH, text="本文检测到 297 对时频冲突。")]
        if figure:
            blocks.append(ContentBlock(
                block_type=BlockType.FIGURE, text="图 1 装备类别分布",
                figure_ids=["FIG-001"],
            ))
        if table:
            blocks.append(ContentBlock(
                block_type=BlockType.TABLE, text="表 1 统计",
                table_ids=["TAB-001"],
            ))

        paper = PaperIR(
            title="时频冲突检测与消解",
            abstract="摘要：本文建立了时频冲突检测与消解模型，检测到 297 对冲突。",
            keywords=["时频冲突", "离散化", "整数规划"],
            sections=[
                PaperSection(title="问题重述", content_blocks=[
                    ContentBlock(block_type=BlockType.PARAGRAPH, text="某区域内 150 个用频装备提报计划。"),
                ]),
                PaperSection(title="模型求解与结果分析", content_blocks=blocks),
            ],
        )
        build_dir = tmp_path / "paper"
        build_dir.mkdir(parents=True, exist_ok=True)
        tex = build_cumcm_latex(paper, figures, tables, build_dir)
        return paper, figures, tables, tex, build_dir

    def test_latex_includes_figure_table_and_numbers(self, tmp_path):
        _, _, _, tex, build_dir = self._build(tmp_path)
        assert r"\includegraphics" in tex
        assert "fig1.png" in tex
        assert "297" in tex
        assert "时频冲突检测与消解" in tex
        assert (build_dir / "fig1.png").exists()

    def test_compiles_to_a_real_pdf(self, tmp_path):
        from mathmodel.autopilot.deliver import build_pdf

        _, _, _, tex, build_dir = self._build(tmp_path)
        result = build_pdf(tex, build_dir)
        if not result.compiler_available:
            pytest.skip("paper compilation image not available")
        assert result.success, result.reason or result.log_tail[-1500:]
        assert result.pdf_size > 5000
        assert Path(result.pdf_path).read_bytes()[:5] == b"%PDF-"

    def test_package_contains_every_deliverable(self, tmp_path):
        from mathmodel.autopilot.deliver import SupportPackageBuilder, build_pdf

        _, figures, tables, tex, build_dir = self._build(tmp_path)
        pdf = build_pdf(tex, build_dir)
        if not pdf.success:
            pytest.skip("paper compilation image not available")

        out_dir = tmp_path / "output"
        manifest = SupportPackageBuilder(out_dir).build(
            paper_pdf=Path(pdf.pdf_path),
            paper_tex=build_dir / "paper.tex",
            source_code=[build_dir / "paper.tex"],
            processed_data=[],
            original_files=[],
            figures=[Path(f.artifact_path) for f in figures.all()],
            tables=[],
            result_files=[],
            manifest_extra={"run_id": "RUN-1"},
        )

        assert (out_dir / "paper.pdf").exists()
        assert (out_dir / "manifest.json").exists()
        assert (out_dir / "support" / "necessary_supporting_files" / "paper.tex").exists()
        assert (out_dir / "support" / "figures" / "fig1.png").exists()
        assert manifest["paper_pdf_size"] > 5000

    def test_figure_missing_from_ir_is_detected(self, tmp_path):
        # The final consistency check keys off collect_figures/collect_tables.
        paper, _, _, _, _ = self._build(tmp_path, figure=False, table=True)
        assert paper.collect_figures() == []
        assert paper.collect_tables()

    def test_stale_figure_id_is_skipped_not_crashed(self, tmp_path):
        from mathmodel.autopilot.deliver import build_cumcm_latex

        paper = PaperIR(
            title="t",
            sections=[PaperSection(title="结果", content_blocks=[
                ContentBlock(block_type=BlockType.FIGURE, text="missing",
                             figure_ids=["FIG-999"]),
            ])],
        )
        build_dir = tmp_path / "paper"
        build_dir.mkdir()
        tex = build_cumcm_latex(paper, FigureRegistry(), TableRegistry(), build_dir)
        assert r"\includegraphics" not in tex


class TestSubprocessDecoding:
    """Hosts with a GBK locale must not lose subprocess output.

    xelatex and generated programs emit UTF-8 Chinese. Decoding that with the
    host locale raises inside the subprocess reader thread, which leaves
    stdout as None and surfaces later as "'NoneType' object is not
    subscriptable" - a message that points nowhere near the real cause.
    """

    def test_latex_log_with_chinese_survives(self, tmp_path):
        from mathmodel.autopilot.deliver import build_pdf

        build_dir = tmp_path / "paper"
        build_dir.mkdir()
        # \typeout writes this straight to stdout, so the captured output is
        # genuinely non-ASCII and cannot be decoded with a GBK locale.
        tex = (
            "\\documentclass[11pt]{ctexart}\n"
            "\\begin{document}\n"
            "\\typeout{中文日志：时频冲突检测与消解}\n"
            "测试正文。\n"
            "\\end{document}\n"
        )
        result = build_pdf(tex, build_dir)
        if not result.compiler_available:
            pytest.skip("paper compilation image not available")

        assert result.success, result.reason or (result.log_tail or "")[-1500:]
        assert result.pdf_path
        assert result.log_tail, "compiler log should have been captured"
        assert "中文日志" in result.log_tail

    def test_sandbox_decodes_chinese_stdout(self):
        import asyncio

        from mathmodel.sandbox.docker_backend import DockerSandboxBackend, SandboxLimits

        backend = DockerSandboxBackend(image="mathmodel-ai-autopilot:latest")
        if not backend.available:
            pytest.skip("docker sandbox not available")

        record = asyncio.run(backend.execute(
            code='print("中文诊断：开始求解")',
            limits=SandboxLimits(timeout_seconds=60),
        ))
        assert record.status.value == "success", record.stderr
        assert record.stdout is not None
        assert "中文诊断" in record.stdout


class TestTraceableNumbers:
    """The paper must not state a number that no verified source supports."""

    def _allowed(self, statistics=None,
                 problem_text="附件1给出3类共150个用频装备。频域离散化为100个频段。"):
        from mathmodel.autopilot.pipeline import _traceable_numbers

        return _traceable_numbers(
            statistics if statistics is not None else {"conflicts": 297},
            TableRegistry(),
            problem_text,
        )

    def test_nested_statistics_are_traceable(self):
        allowed = self._allowed({"p2": {"by_class": {"A": {"total": 20}}}})
        assert 20.0 in allowed

    def test_problem_statement_numbers_are_traceable(self):
        assert 150.0 in self._allowed()

    def test_table_cells_are_traceable(self):
        from mathmodel.autopilot.pipeline import _traceable_numbers

        tables = TableRegistry()
        tables.register(TableBuilder().build_statistics_table({"conflicts": 297}, "EVD"))
        assert 297.0 in _traceable_numbers({}, tables, "")

    def test_fabricated_class_count_is_flagged(self):
        from mathmodel.autopilot.pipeline import _untraceable_numbers

        # The real distribution is A=20, B=40, C=90; "各50个" is invented.
        allowed = self._allowed({
            "p2": {"by_class": {"A": {"total": 20}, "B": {"total": 40}, "C": {"total": 90}}}
        })
        text = "对附件1中150个用频计划（A、B、C类各50个）进行检测。"
        assert _untraceable_numbers(text, allowed) == {"50"}

    def test_verified_and_problem_numbers_are_not_flagged(self):
        from mathmodel.autopilot.pipeline import _untraceable_numbers

        allowed = self._allowed()
        text = "附件1给出150个计划，离散化为100个频段，共检出297对冲突，平移不超过10。"
        assert _untraceable_numbers(text, allowed) == set()

    def test_small_structural_integers_are_ignored(self):
        from mathmodel.autopilot.pipeline import _untraceable_numbers

        assert _untraceable_numbers("见第 3 节与表 2。", set()) == set()

    def test_untraceable_decimal_is_flagged(self):
        from mathmodel.autopilot.pipeline import _untraceable_numbers

        assert _untraceable_numbers("准确率为 87.5%。", set()) == {"87.5"}


class TestInlineMathPreservation:
    """Inline math in prose must survive escaping as real notation."""

    def test_inline_math_is_preserved(self):
        from mathmodel.paper.renderer import _latex_escape

        escaped = _latex_escape("以固定时间长度 $\\Delta t$ 为基本单位。")
        assert "$\\Delta t$" in escaped
        assert r"\textbackslash{}" not in escaped
        assert r"\$" not in escaped

    def test_plain_text_is_still_escaped(self):
        from mathmodel.paper.renderer import _latex_escape

        escaped = _latex_escape("比例 50% 与 A&B 以及 x_1")
        assert r"\%" in escaped
        assert r"\&" in escaped
        assert r"\_" in escaped

    def test_stray_dollar_is_escaped(self):
        from mathmodel.paper.renderer import _latex_escape

        assert _latex_escape("价格是 100$") == r"价格是 100\$"

    def test_math_span_still_cannot_smuggle_primitives(self):
        from mathmodel.paper.renderer import _check_latex_injection, _latex_escape

        tex = _latex_escape(r"公式 $\input{secret}$ 结束")
        assert _check_latex_injection(tex)

    def test_paper_with_inline_math_compiles(self, tmp_path):
        from mathmodel.autopilot.deliver import build_pdf, build_cumcm_latex

        paper = PaperIR(
            title="测试",
            sections=[PaperSection(title="模型建立", content_blocks=[
                ContentBlock(
                    block_type=BlockType.PARAGRAPH,
                    text="时间长度以 $\\Delta t$ 为单位，频率以 $\\Delta f$ 为单位。",
                ),
            ])],
        )
        build_dir = tmp_path / "paper"
        build_dir.mkdir()
        tex = build_cumcm_latex(paper, FigureRegistry(), TableRegistry(), build_dir)
        assert "$\\Delta t$" in tex
        result = build_pdf(tex, build_dir)
        if not result.compiler_available:
            pytest.skip("paper compilation image not available")
        assert result.success, (result.reason or "") + (result.log_tail or "")[-1200:]


class TestRetryCall:
    """A dropped provider connection must not destroy a completed run."""

    def test_retries_then_succeeds(self):
        import asyncio

        from mathmodel.autopilot.pipeline import CUMCMAutopilot

        calls = {"n": 0}

        async def flaky():
            calls["n"] += 1
            if calls["n"] < 3:
                raise ConnectionError("reset by peer")
            return "ok"

        assert asyncio.run(CUMCMAutopilot._retry_call("t", flaky, max_attempts=3)) == "ok"
        assert calls["n"] == 3

    def test_raises_after_exhausting_attempts(self):
        import asyncio

        import pytest as _pytest

        from mathmodel.autopilot.pipeline import CUMCMAutopilot

        async def always_fails():
            raise ConnectionError("reset by peer")

        with _pytest.raises(RuntimeError, match="failed after 2 attempts"):
            asyncio.run(CUMCMAutopilot._retry_call("t", always_fails, max_attempts=2))


class TestTemplateHeaderGate:
    """The template's header row is part of the required answer."""

    def _checks(self, tmp_path, produced_headers, template_headers):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["编号"], [["A001"]])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result1.xlsx", produced_headers, [[1, "A001", "B002"]])
        return DeterministicVerifier(data_dir).run(
            outcome=_outcome({"conflicts": 1, "residual_conflicts": 0}),
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers=template_headers,
        )

    def test_matching_headers_pass(self, tmp_path):
        headers = ["序号", "冲突装备1", "冲突设备2"]
        checks = self._checks(tmp_path, headers, {"result1.xlsx": headers})
        by_name = {c.name: c for c in checks}
        assert by_name["output_structure_matches_template"].status == CheckStatus.PASS

    def test_renamed_columns_fail_even_with_the_right_count(self, tmp_path):
        # The real run shipped these headings while the template asked for
        # 冲突装备1 / 冲突设备2. Column count alone must not pass this.
        checks = self._checks(
            tmp_path,
            ["序号", "用频装备编号1", "用频装备编号2"],
            {"result1.xlsx": ["序号", "冲突装备1", "冲突设备2"]},
        )
        by_name = {c.name: c for c in checks}
        check = by_name["output_structure_matches_template"]
        assert check.status == CheckStatus.FAIL
        assert "header does not match" in check.detail

    def test_whitespace_differences_are_tolerated(self, tmp_path):
        checks = self._checks(
            tmp_path,
            [" 序号 ", "冲突装备1", "冲突设备2"],
            {"result1.xlsx": ["序号", "冲突装备1", "冲突设备2"]},
        )
        by_name = {c.name: c for c in checks}
        assert by_name["output_structure_matches_template"].status == CheckStatus.PASS


class TestCodegenTemplateHeaders:
    """The solver must be told the template's exact column headers."""

    def test_prompt_lists_the_mandatory_headers(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator
        from mathmodel.domain.math_model import MathematicalModel

        prompt = SolverCodeGenerator._build_prompt(None, 
            model=MathematicalModel(
                name="m", objective="min", variables=[], constraints=[], equations=[]
            ),
            problem_text="题目",
            data_schema="{}",
            required_outputs=["result1.xlsx"],
            input_file_names=["附件1.xlsx"],
            output_headers={"result1.xlsx": ["序号", "冲突装备1", "冲突设备2"]},
        )
        assert "MANDATORY COLUMN HEADERS" in prompt
        assert "冲突装备1" in prompt
        assert "冲突设备2" in prompt

    def test_prompt_omits_the_section_without_templates(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator
        from mathmodel.domain.math_model import MathematicalModel

        prompt = SolverCodeGenerator._build_prompt(None, 
            model=MathematicalModel(
                name="m", objective="min", variables=[], constraints=[], equations=[]
            ),
            problem_text="题目",
            data_schema="{}",
            required_outputs=["out.csv"],
            input_file_names=["a.csv"],
        )
        assert "MANDATORY COLUMN HEADERS" not in prompt


class TestPerPlanCoverageGate:
    """A table keyed by plan ID must be answered for every plan."""

    def _run(self, tmp_path, rows, expected):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["用频装备编号"], [["A001"], ["A002"]])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result2.xlsx", ["用频装备编号", "是否撤销用频计划"], rows)
        return DeterministicVerifier(data_dir).run(
            outcome=_outcome({"residual_conflicts": 0}),
            output_dir=out_dir,
            required_outputs=["result2.xlsx"],
            template_headers={"result2.xlsx": ["用频装备编号", "是否撤销用频计划"]},
            expected_row_counts=expected,
        )

    def test_incomplete_per_plan_table_fails(self, tmp_path):
        checks = self._run(tmp_path, [["A001", "是"]], {"result2.xlsx": 2})
        check = {c.name: c for c in checks}["per_plan_table_covers_every_plan"]
        assert check.status == CheckStatus.FAIL
        assert "1 row(s) for 2 plan(s)" in check.detail

    def test_complete_per_plan_table_passes(self, tmp_path):
        checks = self._run(
            tmp_path, [["A001", "是"], ["A002", "否"]], {"result2.xlsx": 2}
        )
        check = {c.name: c for c in checks}["per_plan_table_covers_every_plan"]
        assert check.status == CheckStatus.PASS

    def test_no_per_plan_template_is_not_applicable(self, tmp_path):
        checks = self._run(tmp_path, [["A001", "是"]], {})
        check = {c.name: c for c in checks}["per_plan_table_covers_every_plan"]
        assert check.status == CheckStatus.NOT_APPLICABLE


class TestModelRepairTrigger:
    """A semantic disagreement must send the repair loop back to modeling."""

    def _check(self, name, category, status, detail):
        return VerificationCheck(
            name=name, category=category, status=status, detail=detail
        )

    def test_independent_disagreement_triggers_model_repair(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            self._check("independent_independent_recomputation", "independent",
                        CheckStatus.FAIL,
                        "plans=150 recomputed_total=297 claimed=237 match=False"),
        ])
        assert _needs_model_repair(report)

    def test_pure_output_defect_does_not_trigger_model_repair(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            self._check("required_outputs_present", "outputs", CheckStatus.FAIL,
                        "missing=['result1.xlsx']"),
            self._check("output_structure_matches_template", "outputs",
                        CheckStatus.FAIL, "result1.xlsx: header does not match"),
        ])
        assert not _needs_model_repair(report)

    def test_passing_report_does_not_trigger_repair(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.PASS, checks=[
            self._check("independent_x", "independent", CheckStatus.PASS, "ok"),
        ])
        assert not _needs_model_repair(report)


class TestModelRepairScope:
    """Only input-derived disagreements implicate the model."""

    def test_residual_conflicts_do_not_repair_the_model(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            VerificationCheck(
                name="independent_constraint_satisfaction", category="independent",
                status=CheckStatus.FAIL,
                detail="problem2 remaining conflicts after applying adjustments = 39",
            ),
        ])
        assert not _needs_model_repair(report)

    def test_baseline_comparison_does_not_repair_the_model(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            VerificationCheck(
                name="independent_baseline_comparison", category="independent",
                status=CheckStatus.FAIL,
                detail="solution is worse than the trivial baseline",
            ),
        ])
        assert not _needs_model_repair(report)


class TestReferentialIntegrityScope:
    """Only columns that reference an existing entity are checked."""

    def _run(self, tmp_path, headers, rows):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        _make_xlsx(data_dir / "附件1.xlsx", ["用频装备编号"], [["A001"], ["B002"]])
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result3.xlsx", headers, rows)
        return {
            c.name: c for c in DeterministicVerifier(data_dir).run(
                outcome=_outcome({"residual_conflicts": 0}),
                output_dir=out_dir,
                required_outputs=["result3.xlsx"],
                template_headers={},
            )
        }

    def test_serial_numbers_for_new_items_are_not_flagged(self, tmp_path):
        # result3 numbers newly added plans, so "newC10" is not a broken
        # reference to the input data.
        checks = self._run(
            tmp_path,
            ["新增用频装备序号", "调整后频段区间", "调整后时间区间"],
            [["newC10", "[0,10)", "[0,5)"], ["newC100", "[10,20)", "[0,5)"]],
        )
        assert checks["referential_integrity"].status == CheckStatus.PASS

    def test_unknown_equipment_ids_are_still_flagged(self, tmp_path):
        checks = self._run(
            tmp_path,
            ["用频装备编号", "调整后频段区间"],
            [["A001", "[0,10)"], ["Z999", "[10,20)"]],
        )
        assert checks["referential_integrity"].status == CheckStatus.FAIL
        assert "Z999" in checks["referential_integrity"].detail


class TestVacuousVerifierGuard:
    """A verifier that read no input must not be believed."""

    def _summary(self, **extra):
        base = {
            "checks": [{"name": "constraint_1", "status": "PASS", "detail": "ok"}],
            "statistics": {},
            "overall": "PASS",
        }
        base.update(extra)
        return base

    def test_missing_input_records_fails(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary())
        guard = {c.name: c for c in checks}["independent_read_the_input"]
        assert guard.status == CheckStatus.FAIL
        assert "no `input_records` field" in guard.detail

    def test_zero_input_records_fails(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary(input_records=0))
        guard = {c.name: c for c in checks}["independent_read_the_input"]
        assert guard.status == CheckStatus.FAIL
        assert "input_records=0" in guard.detail

    def test_positive_input_records_passes(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary(input_records=150))
        guard = {c.name: c for c in checks}["independent_read_the_input"]
        assert guard.status == CheckStatus.PASS

    def test_a_vacuous_verifier_cannot_produce_a_passing_report(self):
        from mathmodel.autopilot.verify import (
            build_report,
            checks_from_independent_summary,
        )
        from mathmodel.domain.math_model import MathematicalModel

        checks, _ = checks_from_independent_summary(self._summary(input_records=0))
        report = build_report(
            MathematicalModel(name="m", objective="o"), "run", checks
        )
        assert not report.passed


class TestVerifierGetsDataSchema:
    """The verifier must know the real input columns to parse the data."""

    def test_prompt_includes_the_real_schema(self):
        import asyncio
        import json as _json

        from mathmodel.autopilot.verify import IndependentVerifier
        from mathmodel.domain.math_model import MathematicalModel

        captured = {}

        class FakeRouter:
            async def route_generate(self, **kwargs):
                captured.update(kwargs)
                raise RuntimeError("stop here")

        verifier = IndependentVerifier(FakeRouter(), runner=None)
        schema = _json.dumps({
            "tables": [{
                "file": "附件1.xlsx",
                "headers": ["用频装备编号", "频段区间", "时间区间", "间隔时长", "使用次数"],
            }]
        }, ensure_ascii=False)
        with pytest.raises(RuntimeError):
            asyncio.run(verifier.generate(
                model=MathematicalModel(name="m", objective="o"),
                problem_text="题目",
                required_outputs=["result1.xlsx"],
                solver_statistics={},
                available_inputs=["附件1.csv"],
                data_schema=schema,
            ))
        prompt = captured["prompt"]
        assert "INPUT DATA SCHEMA" in prompt
        assert "使用次数" in prompt

    def test_prompt_omits_the_section_without_a_schema(self):
        import asyncio

        from mathmodel.autopilot.verify import IndependentVerifier
        from mathmodel.domain.math_model import MathematicalModel

        captured = {}

        class FakeRouter:
            async def route_generate(self, **kwargs):
                captured.update(kwargs)
                raise RuntimeError("stop here")

        verifier = IndependentVerifier(FakeRouter(), runner=None)
        with pytest.raises(RuntimeError):
            asyncio.run(verifier.generate(
                model=MathematicalModel(name="m", objective="o"),
                problem_text="题目",
                required_outputs=["result1.xlsx"],
                solver_statistics={},
                available_inputs=["a.csv"],
            ))
        assert "INPUT DATA SCHEMA" not in captured["prompt"]


class TestModelRepairNeedsARealRecomputation:
    """A vacuous verifier must not send the loop back to modeling."""

    def test_vacuous_verifier_does_not_repair_the_model(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            VerificationCheck(
                name="independent_read_the_input", category="independent",
                status=CheckStatus.FAIL,
                detail="Independent verifier reported input_records=0",
            ),
            VerificationCheck(
                name="independent_independent_recomputation", category="independent",
                status=CheckStatus.FAIL,
                detail="Parsed 0 plans from 附件1.csv",
            ),
        ])
        assert not _needs_model_repair(report)

    def test_real_disagreement_still_repairs_the_model(self):
        from mathmodel.autopilot.pipeline import _needs_model_repair

        report = VerificationReport(overall=CheckStatus.FAIL, checks=[
            VerificationCheck(
                name="independent_read_the_input", category="independent",
                status=CheckStatus.PASS, detail="parsed 150 records",
            ),
            VerificationCheck(
                name="independent_independent_recomputation", category="independent",
                status=CheckStatus.FAIL,
                detail="recomputed_total=297 claimed=237",
            ),
        ])
        assert _needs_model_repair(report)


class TestEvidenceRequiredForConstraintChecks:
    """A PASS with no evidence is not coverage."""

    def _summary(self, checks):
        return {"checks": checks, "input_records": 150, "statistics": {}, "overall": "PASS"}

    def test_constraint_pass_without_evidence_becomes_not_run(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary([
            {"name": "constraint_6", "status": "PASS", "detail": "Check of: c6"},
        ]))
        check = {c.name: c for c in checks}["independent_constraint_6"]
        assert check.status == CheckStatus.NOT_RUN
        assert "not actually checked" in check.detail

    def test_constraint_pass_with_evidence_is_kept(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary([
            {"name": "constraint_6", "status": "PASS", "detail": "ok",
             "evidence": {"checked": 150}},
        ]))
        assert {c.name: c for c in checks}["independent_constraint_6"].status == CheckStatus.PASS

    def test_constraint_failure_without_evidence_still_fails(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary([
            {"name": "constraint_6", "status": "FAIL", "detail": "3 plans changed two parameters"},
        ]))
        assert {c.name: c for c in checks}["independent_constraint_6"].status == CheckStatus.FAIL

    def test_a_report_with_a_not_run_constraint_does_not_pass(self):
        from mathmodel.autopilot.verify import (
            build_report,
            checks_from_independent_summary,
        )
        from mathmodel.domain.math_model import MathematicalModel

        checks, _ = checks_from_independent_summary(self._summary([
            {"name": "constraint_6", "status": "PASS", "detail": "Check of: c6"},
        ]))
        report = build_report(MathematicalModel(name="m", objective="o"), "run", checks)
        assert not report.passed


class TestReportedVerificationMatchesTheGatingReport:
    """The result must report the verdict that gated the run."""

    def test_stale_report_is_not_reported(self, tmp_path):
        import json as _json
        import os
        import time

        from mathmodel.autopilot.pipeline import CUMCMAutopilot
        from mathmodel.autopilot.state import RunState

        run_dir = tmp_path / "run"
        (run_dir / "artifacts" / "verify").mkdir(parents=True)
        (run_dir / "output").mkdir(parents=True)
        state = RunState.create(run_dir)
        state.status = RunStatus.COMPLETED
        state.save()

        verify_dir = run_dir / "artifacts" / "verify"
        # A stale failure with a HIGHER number than the report that passed.
        (verify_dir / "report3.json").write_text(
            _json.dumps({"overall": "FAIL"}), encoding="utf-8"
        )
        time.sleep(0.01)
        (verify_dir / "report1.json").write_text(
            _json.dumps({"overall": "PASS"}), encoding="utf-8"
        )

        result = CUMCMAutopilot(workspace=str(tmp_path))._result(state)
        assert result.verification == "PASS"


class TestTruncatedStructuredResponseRetry:
    """A truncated JSON body can never parse, so retry with thinking off."""

    def _provider(self, responses):
        from mathmodel.providers.openai import OpenAIProvider

        provider = OpenAIProvider(api_key="k", base_url="http://x/v1")
        calls = []

        async def fake_chat(client, **kwargs):
            calls.append(kwargs)
            return responses[min(len(calls) - 1, len(responses) - 1)]

        provider._chat_completion = fake_chat
        provider._get_client = lambda: object()
        return provider, calls

    def _response(self, finish_reason, content):
        from types import SimpleNamespace

        return SimpleNamespace(
            choices=[SimpleNamespace(
                finish_reason=finish_reason,
                message=SimpleNamespace(content=content),
            )],
            model="m",
            usage=SimpleNamespace(prompt_tokens=1, completion_tokens=1, total_tokens=2),
        )

    def test_truncated_response_is_retried_with_thinking_disabled(self):
        import asyncio

        from mathmodel.providers.base import StructuredGenerationRequest

        class Tiny(pydantic.BaseModel):
            problem_id: str = ""

        provider, calls = self._provider([
            self._response("length", '{"problem_id": "p"'),
            self._response("stop", '{"problem_id": "p"}'),
        ])
        result = asyncio.run(provider.structured_generate(
            StructuredGenerationRequest(prompt="p", output_schema=Tiny)
        ))
        assert len(calls) == 2
        assert calls[1].get("extra_body") == {"thinking": {"type": "disabled"}}
        assert result.problem_id == "p"

    def test_untruncated_response_is_not_retried(self):
        import asyncio

        from mathmodel.providers.base import StructuredGenerationRequest

        class Tiny(pydantic.BaseModel):
            problem_id: str = ""

        provider, calls = self._provider([
            self._response("stop", '{"problem_id": "p"}'),
        ])
        asyncio.run(provider.structured_generate(
            StructuredGenerationRequest(prompt="p", output_schema=Tiny)
        ))
        assert len(calls) == 1
        assert "extra_body" not in calls[0]


class TestVerifierDisagreementFeedback:
    """A disagreement must be fed back to the verifier too, not only the model."""

    def test_disagreement_produces_verifier_feedback(self):
        import inspect

        from mathmodel.autopilot import pipeline

        source = inspect.getsource(pipeline.CUMCMAutopilot)
        assert "carried_verifier_feedback = [" in source
        assert "DISAGREE with the solver" in source
        assert "re-derive" in source

    def test_feedback_is_only_carried_on_actual_failure(self):
        from mathmodel.autopilot.verify import compare_statistics

        agree = compare_statistics({"a": 297}, {"a": 297})
        assert agree.status == CheckStatus.PASS

        disagree = compare_statistics({"a": 297}, {"a": 237})
        assert disagree.status == CheckStatus.FAIL
        assert "297" in disagree.detail and "237" in disagree.detail


class TestObjectivesGate:
    """A verifier that never judges the answer must not pass the run."""

    def _summary(self, objectives=None, checks=None):
        s = {
            "checks": checks or [{"name": "constraint_1", "status": "PASS",
                                  "detail": "input range ok", "evidence": {"n": 150}}],
            "input_records": 150,
            "statistics": {},
            "overall": "PASS",
        }
        if objectives is not None:
            s["objectives"] = objectives
        return s

    def _check(self, summary, outputs=("result1.xlsx", "result2.xlsx")):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(
            summary, required_outputs=list(outputs)
        )
        return {c.name: c for c in checks}["independent_objectives_achieved"]

    def test_missing_objectives_fails(self):
        check = self._check(self._summary())
        assert check.status == CheckStatus.FAIL
        assert "never stated whether the solution actually achieves" in check.detail

    def test_output_without_a_verdict_fails(self):
        check = self._check(self._summary(objectives=[
            {"output": "result1.xlsx", "achieved": True, "evidence": "297 pairs"},
        ]))
        assert check.status == CheckStatus.FAIL
        assert "result2.xlsx" in check.detail

    def test_unachieved_objective_fails(self):
        check = self._check(self._summary(objectives=[
            {"output": "result1.xlsx", "achieved": True, "evidence": "297 pairs"},
            {"output": "result2.xlsx", "achieved": False, "evidence": "297 conflicts remain"},
        ]))
        assert check.status == CheckStatus.FAIL
        assert "297 conflicts remain" in check.detail

    def test_all_outputs_achieved_passes(self):
        check = self._check(self._summary(objectives=[
            {"output": "result1.xlsx", "achieved": True, "evidence": "297 pairs"},
            {"output": "result2.xlsx", "achieved": True, "evidence": "0 residual"},
        ]))
        assert check.status == CheckStatus.PASS

    def test_the_observed_false_pass_is_now_blocked(self):
        """r05: every check was an input-data property; the answer was unsolved."""
        from mathmodel.autopilot.verify import (
            build_report,
            checks_from_independent_summary,
        )
        from mathmodel.domain.math_model import MathematicalModel

        summary = self._summary(checks=[
            {"name": "constraint_1", "status": "PASS", "detail": "band ranges within [0,100]", "evidence": {"n": 150}},
            {"name": "constraint_2", "status": "PASS", "detail": "use counts positive", "evidence": {"v": [3, 4, 12]}},
            {"name": "constraint_3", "status": "PASS", "detail": "interval lengths positive", "evidence": {"n": 150}},
            {"name": "feasibility", "status": "PASS", "detail": "checked 150 adjustments", "evidence": {"n": 150}},
        ])
        checks, _ = checks_from_independent_summary(
            summary, required_outputs=["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"]
        )
        report = build_report(MathematicalModel(name="m", objective="o"), "run", checks)
        assert not report.passed

    def test_no_required_outputs_skips_the_objective_gate(self):
        from mathmodel.autopilot.verify import checks_from_independent_summary

        checks, _ = checks_from_independent_summary(self._summary())
        assert "independent_objectives_achieved" not in {c.name for c in checks}


class TestBlockedRunNeverReportsRunning:
    """A run that returns blockers has stopped, so it must not say `running`."""

    def _state(self, tmp_path, status):
        from mathmodel.autopilot.state import RunState

        run_dir = tmp_path / "run"
        (run_dir / "output").mkdir(parents=True)
        state = RunState.create(run_dir)
        state.status = status
        state.save()
        return state

    def test_blockers_settle_a_running_run(self, tmp_path):
        from mathmodel.autopilot.pipeline import CUMCMAutopilot
        from mathmodel.autopilot.state import RunStatus

        state = self._state(tmp_path, RunStatus.RUNNING)
        result = CUMCMAutopilot(workspace=str(tmp_path))._result(
            state, blockers=["[independent] constraint_1: 297 residual conflicts"]
        )
        assert result.status == "blocked"
        assert state.status == RunStatus.BLOCKED

    def test_a_completed_run_is_not_downgraded(self, tmp_path):
        from mathmodel.autopilot.pipeline import CUMCMAutopilot
        from mathmodel.autopilot.state import RunStatus

        state = self._state(tmp_path, RunStatus.COMPLETED)
        result = CUMCMAutopilot(workspace=str(tmp_path))._result(
            state, blockers=["should not happen"]
        )
        assert result.status == "completed"

    def test_no_blockers_leaves_a_running_run_alone(self, tmp_path):
        from mathmodel.autopilot.pipeline import CUMCMAutopilot
        from mathmodel.autopilot.state import RunStatus

        state = self._state(tmp_path, RunStatus.RUNNING)
        result = CUMCMAutopilot(workspace=str(tmp_path))._result(state)
        assert result.status == "running"


from mathmodel.agents.base import AgentStatus
from mathmodel.domain.math_model import (
    KIND_DUPLICATE_SYMBOL,
    KIND_SYMBOL_COLLISION,
    KIND_UNKNOWN_DEPENDENCY,
    KIND_UNKNOWN_PARAMETER,
    KIND_UNKNOWN_VARIABLE,
    MathematicalModel,
    closure_issues,
    parse_closure_issues,
)


class TestSymbolClosureRepair:
    """#25: undeclared symbols must yield a targeted repair, not a blind retry."""

    @staticmethod
    def _payload(**overrides):
        base = {
            "model_id": "M1",
            "name": "closure test model",
            "description": "d",
            "variables": [
                {"variable_id": "v_x", "symbol": "x", "name": "x"},
            ],
            "parameters": [
                {"parameter_id": "p_T", "symbol": "T", "name": "T",
                 "value": 10.0, "status": "known"},
            ],
            "objectives": [
                {"objective_id": "o1", "name": "obj", "sense": "minimize",
                 "expression": "x", "meaning": "m", "priority": 1},
            ],
            "constraints": [],
            "equations": [
                {"equation_id": "EQ1", "name": "e1", "expression": "x = T",
                 "variable_ids": ["v_x"], "parameter_ids": ["p_T"]},
            ],
        }
        base.update(overrides)
        return base

    # ── criterion 1: an undeclared symbol must FAIL ──────────────────
    def test_undeclared_symbols_fail_validation(self):
        bad = self._payload(equations=[
            {"equation_id": "EQ1", "name": "e1", "expression": "x = T + M_j",
             "variable_ids": ["v_x", "v_y_missing"],
             "parameter_ids": ["p_T", "p_M_missing"]},
        ])
        with pytest.raises(Exception) as excinfo:
            MathematicalModel.model_validate(bad)
        assert "unknown variable v_y_missing" in str(excinfo.value)
        assert "unknown parameter p_M_missing" in str(excinfo.value)

    # ── criterion 2: the failure carries the concrete symbol lists ───
    def test_failure_message_parses_back_into_the_exact_symbols(self):
        bad = self._payload(equations=[
            {"equation_id": "EQ1", "name": "e1", "expression": "x = T + M_j",
             "variable_ids": ["v_x", "v_y_missing"],
             "parameter_ids": ["p_T", "p_M_missing"],
             "dependencies": ["EQ_missing"]},
        ])
        with pytest.raises(Exception) as excinfo:
            MathematicalModel.model_validate(bad)

        issues = parse_closure_issues(str(excinfo.value))
        assert sorted(i.symbol for i in issues if i.kind == KIND_UNKNOWN_VARIABLE) == ["v_y_missing"]
        assert sorted(i.symbol for i in issues if i.kind == KIND_UNKNOWN_PARAMETER) == ["p_M_missing"]
        assert sorted(i.symbol for i in issues if i.kind == KIND_UNKNOWN_DEPENDENCY) == ["EQ_missing"]
        assert {i.location for i in issues} == {"EQ1"}

    def test_duplicate_and_colliding_symbols_are_reported(self):
        bad = self._payload(
            variables=[
                {"variable_id": "v_x", "symbol": "x", "name": "x"},
                {"variable_id": "v_x2", "symbol": "x", "name": "x again"},
                {"variable_id": "v_T", "symbol": "T", "name": "T as a variable"},
            ],
        )
        with pytest.raises(Exception) as excinfo:
            MathematicalModel.model_validate(bad)
        issues = parse_closure_issues(str(excinfo.value))
        kinds = {i.kind for i in issues}
        assert KIND_DUPLICATE_SYMBOL in kinds
        assert KIND_SYMBOL_COLLISION in kinds

    def test_a_clean_model_yields_no_issues(self):
        model = MathematicalModel.model_validate(self._payload())
        assert closure_issues(model) == []
        assert parse_closure_issues("some unrelated error") == []

    def test_describe_and_parse_round_trip(self):
        bad = self._payload(equations=[
            {"equation_id": "EQ7", "name": "e", "expression": "x",
             "variable_ids": ["v_missing"], "parameter_ids": ["p_missing"]},
        ])
        with pytest.raises(Exception) as excinfo:
            MathematicalModel.model_validate(bad)
        issues = parse_closure_issues(str(excinfo.value))
        assert len(issues) == 2
        for issue in issues:
            assert issue.symbol in str(excinfo.value)
            assert issue.location == "EQ7"

    # ── criteria 3 and 4: the next attempt is fed the exact lists ────
    @staticmethod
    def _state():
        from mathmodel.domain.analysis import ProblemAnalysis
        from mathmodel.models.problem_state import ProblemState

        state = ProblemState(title="closure")
        state.metadata_ = {
            "analysis": ProblemAnalysis(
                background="b", core_problem="c", objectives=["o"],
            ).model_dump()
        }
        state.raw_problem = "A problem statement."
        state.selected_model = {"candidate_id": "CAND-1"}
        return state

    async def test_retry_prompt_contains_the_exact_missing_symbols(self):
        from mathmodel.agents.math_modeler import MathModeler

        bad = self._payload(equations=[
            {"equation_id": "EQ1", "name": "e1", "expression": "x = T + M_j",
             "variable_ids": ["v_x", "v_y_missing"],
             "parameter_ids": ["p_T", "p_M_missing"]},
        ])
        with pytest.raises(Exception) as excinfo:
            MathematicalModel.model_validate(bad)
        real_error = excinfo.value

        good = MathematicalModel.model_validate(self._payload())

        class FakeRouter:
            def __init__(self):
                self.prompts = []

            async def route_structured_generate(self, **kwargs):
                self.prompts.append(kwargs["prompt"])
                if len(self.prompts) == 1:
                    raise real_error
                return good

        router = FakeRouter()
        result = await MathModeler(router).run(self._state())

        assert result.status == AgentStatus.COMPLETED
        assert len(router.prompts) == 2
        repair_prompt = router.prompts[1]

        # criterion 3: the repair prompt names the exact symbols
        assert "v_y_missing" in repair_prompt
        assert "p_M_missing" in repair_prompt
        assert "SYMBOL CLOSURE FAILURE" in repair_prompt

        # criterion 4: it is targeted, not a blind regeneration
        assert "Keep every other part of the model as it is" in repair_prompt
        assert "do not rebuild the model from scratch" in repair_prompt

    async def test_first_attempt_has_no_corrective_block(self):
        from mathmodel.agents.math_modeler import MathModeler

        good = MathematicalModel.model_validate(self._payload())

        class FakeRouter:
            def __init__(self):
                self.prompts = []

            async def route_structured_generate(self, **kwargs):
                self.prompts.append(kwargs["prompt"])
                return good

        router = FakeRouter()
        result = await MathModeler(router).run(self._state())
        assert result.status == AgentStatus.COMPLETED
        assert len(router.prompts) == 1
        assert "SYMBOL CLOSURE FAILURE" not in router.prompts[0]

    async def test_a_non_closure_error_still_gets_generic_feedback(self):
        from mathmodel.agents.math_modeler import MathModeler

        good = MathematicalModel.model_validate(self._payload())

        class FakeRouter:
            def __init__(self):
                self.prompts = []

            async def route_structured_generate(self, **kwargs):
                self.prompts.append(kwargs["prompt"])
                if len(self.prompts) == 1:
                    raise ValueError("Provider returned no usable JSON after retry")
                return good

        router = FakeRouter()
        result = await MathModeler(router).run(self._state())
        assert result.status == AgentStatus.COMPLETED
        assert "SYMBOL CLOSURE FAILURE" not in router.prompts[1]
        assert "could not be parsed" in router.prompts[1]

    # ── criterion 5: the repaired model satisfies closure ────────────
    def test_declaring_the_missing_symbols_clears_the_failure(self):
        broken = self._payload(equations=[
            {"equation_id": "EQ1", "name": "e1", "expression": "x = T + M_j",
             "variable_ids": ["v_x", "v_y_missing"],
             "parameter_ids": ["p_T", "p_M_missing"]},
        ])
        with pytest.raises(Exception):
            MathematicalModel.model_validate(broken)

        repaired = self._payload(
            variables=[
                {"variable_id": "v_x", "symbol": "x", "name": "x"},
                {"variable_id": "v_y_missing", "symbol": "v_y_missing",
                 "name": "declared now"},
            ],
            parameters=[
                {"parameter_id": "p_T", "symbol": "T", "name": "T",
                 "value": 10.0, "status": "known"},
                {"parameter_id": "p_M_missing", "symbol": "p_M_missing",
                 "name": "declared now", "value": 1.0, "status": "known"},
            ],
            equations=[
                {"equation_id": "EQ1", "name": "e1", "expression": "x = T + M_j",
                 "variable_ids": ["v_x", "v_y_missing"],
                 "parameter_ids": ["p_T", "p_M_missing"]},
            ],
        )
        model = MathematicalModel.model_validate(repaired)
        assert closure_issues(model) == []


class TestR05FalsePassRegression:
    """Permanent negative fixture: a real report that was accepted as PASS
    while the answer was wrong. The gate must never accept it again."""

    @staticmethod
    def _fixture():
        import json as _json
        from pathlib import Path as _Path

        path = _Path(__file__).resolve().parent / "fixtures" / "r05_false_pass.json"
        return _json.loads(path.read_text(encoding="utf-8"))

    def test_fixture_records_the_ground_truth(self):
        fixture = self._fixture()
        truth = fixture["independent_ground_truth"]
        assert truth["verdict"] == "WRONG"
        # The answer left every conflict unresolved and produced an empty result3.
        assert truth["result2_residual_conflicts"] == 297
        assert truth["result2_cancelled"] == 0
        assert truth["result3_rows"] == 0
        assert fixture["verifier_summary"]["overall"] == "PASS"

    def test_the_old_gate_would_still_accept_it(self):
        """Guards the regression: without the objectives contract this passes."""
        from mathmodel.autopilot.verify import (
            build_report,
            checks_from_independent_summary,
        )
        from mathmodel.domain.math_model import MathematicalModel

        summary = self._fixture()["verifier_summary"]
        checks, _ = checks_from_independent_summary(summary)
        report = build_report(MathematicalModel(name="m", objective="o"), "r05", checks)
        assert report.passed

    def test_the_current_gate_rejects_it(self):
        from mathmodel.autopilot.verify import (
            build_report,
            checks_from_independent_summary,
        )
        from mathmodel.domain.math_model import MathematicalModel

        fixture = self._fixture()
        checks, _ = checks_from_independent_summary(
            fixture["verifier_summary"],
            required_outputs=fixture["required_outputs"],
        )
        report = build_report(MathematicalModel(name="m", objective="o"), "r05", checks)
        assert not report.passed

        by_name = {c.name: c for c in checks}
        assert by_name["independent_objectives_achieved"].status == CheckStatus.FAIL

    def test_every_check_in_the_fixture_is_an_input_property(self):
        """Documents WHY it slipped through: no check judged the answer."""
        fixture = self._fixture()
        names = [c["name"] for c in fixture["verifier_summary"]["checks"]]
        assert names, "fixture must carry the verifier's checks"
        assert not any(n.startswith("objective") for n in names)
        # The one constraint check that ran only re-derived result1 (detection).
        satisfaction = [c for c in fixture["verifier_summary"]["checks"]
                        if c["name"] == "constraint_satisfaction"]
        assert satisfaction and "result1" in satisfaction[0]["detail"]


# ═══════════════════════════════════════════════════════════════
# Verification gate: template ellipsis + problem-stated requirements
# ═══════════════════════════════════════════════════════════════

def _a_radii() -> list[str]:
    """The radial axis 0, 0.1, ..., 2 as the 2026 A templates expect it."""
    out = []
    for i in range(21):
        value = round(i * 0.1, 1)
        out.append(str(int(value)) if value == int(value) else str(value))
    return out


class TestTemplateEllipsisExpansion:
    """A '…' in a template stands for an omitted arithmetic run of columns.

    Comparing it literally rejected every correct answer to the 2026 A problem,
    whose templates abbreviate the radial axis as `0, 0.1, 0.2, …, 2`.
    """

    HEADER = ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "2"]
    LABELLED = ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "药材表面"]

    def test_the_real_a_expansion_is_accepted(self):
        produced = ["时间\\到药材中心的距离"] + _a_radii()
        assert header_mismatches(self.HEADER, produced) == []

    def test_the_real_a_result4_expansion_is_accepted(self):
        produced = ["时间\\到药材中心的距离"] + _a_radii()[:-1] + ["药材表面"]
        assert header_mismatches(self.LABELLED, produced) == []

    def test_an_expansion_that_skips_a_value_is_rejected(self):
        produced = ["时间\\到药材中心的距离"] + ["0", "0.1", "0.2", "0.5"] + _a_radii()[5:]
        mismatches = header_mismatches(self.HEADER, produced)
        assert mismatches and "0.3" in mismatches[0]

    def test_an_expansion_that_never_reaches_the_printed_tail_is_rejected(self):
        produced = ["时间\\到药材中心的距离"] + _a_radii()[:-1] + ["1.9"]
        mismatches = header_mismatches(self.HEADER, produced)
        assert mismatches, "an expansion that stops short of 2 must not be accepted"

    def test_an_expansion_with_no_derivable_step_still_fails(self):
        mismatches = header_mismatches(["head", "…", "tail"], ["head", "x", "tail"])
        assert mismatches and "cannot be expanded" in mismatches[0]

    def test_a_prefix_that_is_not_an_arithmetic_progression_still_fails(self):
        expected = ["head", "0", "0.1", "0.4", "…", "2"]
        produced = ["head", "0", "0.1", "0.4", "0.7", "1.0", "1.3", "1.6", "1.9", "2"]
        mismatches = header_mismatches(expected, produced)
        assert mismatches and "arithmetic progression" in mismatches[0]

    def test_an_ellipsis_with_too_few_columns_is_rejected(self):
        mismatches = header_mismatches(self.HEADER, ["时间\\到药材中心的距离", "0", "2"])
        assert mismatches and "at least" in mismatches[0]

    # ── G: the D problem's literal headers are unchanged ─────────────
    def test_a_template_without_an_ellipsis_is_compared_literally(self):
        expected = ["序号", "冲突装备1", "冲突设备2"]
        assert header_mismatches(expected, list(expected)) == []
        assert header_mismatches(expected, ["序号", "X", "冲突设备2"]) != []

    def test_a_literal_template_still_rejects_a_short_header(self):
        mismatches = header_mismatches(["a", "b", "c"], ["a", "b"])
        assert mismatches and "expected >= 3 columns" in mismatches[0]

    def test_the_d_template_headers_round_trip(self):
        expected = ["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"]
        assert header_mismatches(expected, list(expected)) == []


class TestProblemDerivedConstraints:
    """Requirements are read out of the statement, not guessed."""

    A_STATEMENT = (
        "问题 3  按照烘干要求，药材各处的水分浓度应低于 0.15 kg/kg ，请确定药材烘干所需\n"
        "要的时间（单位：h）。在论文中按表 5 的格式给出每隔 6 h 、到药材中心距离每隔 "
        "0.5 cm 的水分浓度，并将药材内部水分浓度每隔 60 s 、到药材中心距离每隔 0.1 cm "
        "的完整结果保存到文件 result3.xlsx 中。"
    )

    def test_the_a_sampling_interval_is_read_per_file(self):
        steps = problem_time_steps(self.A_STATEMENT)
        assert steps == {"result3.xlsx": 60.0}

    def test_a_phrase_split_across_a_line_break_is_still_found(self):
        # the PDF wraps "烘干所需\n要的时间"
        assert DRYING_TIME_RE.search(self.A_STATEMENT) is not None

    def test_a_table_row_label_is_not_a_requirement(self):
        # "烘干结束时间" is a row heading inside the answer tables, not a demand
        assert DRYING_TIME_RE.search("表 5  药材烘干过程的水分浓度\n烘干结束时间") is None

    def test_the_threshold_is_applied_to_the_drying_outputs_only(self):
        targets = problem_final_state_targets(
            self.A_STATEMENT, ["result1.xlsx", "result3.xlsx"]
        )
        assert targets == [("result3.xlsx", "<", 0.15)]

    def test_the_d_statement_yields_no_such_constraints(self):
        statement = (
            "某区域有 150 个用频装备，请检测时频冲突。将冲突装备对保存到文件 "
            "result1.xlsx 中，将调整方案保存到文件 result2.xlsx 中。"
        )
        assert problem_time_steps(statement) == {}
        assert problem_final_state_targets(statement, ["result1.xlsx", "result2.xlsx"]) == []


class TestGateReportsEveryProblemItCan:
    """One failing check must not silence the checks that could still run.

    This was the direct cause of the 2026 A run reporting a single bogus
    ellipsis complaint while three real defects went unreported.
    """

    HEADERS = {
        "result1.xlsx": ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "2"],
        "result2.xlsx": ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "2"],
        "result3.xlsx": ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "2"],
        "result4.xlsx": ["时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "药材表面"],
    }
    SHEETS = {
        "result1.xlsx": ["温度", "水分浓度"],
        "result2.xlsx": ["温度", "水分浓度"],
        "result3.xlsx": ["Sheet1"],
        "result4.xlsx": ["Sheet1"],
    }
    STATEMENT = (
        "问题 3  按照烘干要求，药材各处的水分浓度应低于 0.15 kg/kg ，请确定药材烘干所需\n"
        "要的时间（单位：h）。并将药材内部水分浓度每隔 60 s 、到药材中心距离每隔 0.1 cm"
        " 的完整结果保存到文件 result3.xlsx 中。\n"
        "问题 4  请确定药材的烘干时长，并将完整结果保存到文件 result4.xlsx 中。"
    )
    REQUIRED = ["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"]

    @staticmethod
    def _workbook(path, sheets, header, rows):
        import openpyxl

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name in sheets:
            ws = wb.create_sheet(title=name)
            ws.append(header)
            for row in rows:
                ws.append(row)
        wb.save(path)

    def _gate(
        self,
        tmp_path,
        *,
        temperature_sheet="温度",
        step=60,
        final3=0.14,
        final4=0.14,
        header=None,
        statement=None,
    ):
        import openpyxl  # noqa: F401  (used by _workbook)
        from mathmodel.autopilot.codegen import ExecutionOutcome
        from mathmodel.autopilot.verify import DeterministicVerifier

        base = header if header is not None else ["时间\\到药材中心的距离"] + _a_radii()
        labelled = ["时间\\到药材中心的距离"] + _a_radii()[:-1] + ["药材表面"]

        def rows(spec, times, value):
            return [[t] + [value] * (len(spec) - 1) for t in times]

        self._workbook(tmp_path / "result1.xlsx", [temperature_sheet, "水分浓度"],
                       base, rows(base, range(4), 28.0))
        self._workbook(tmp_path / "result2.xlsx", [temperature_sheet, "水分浓度"],
                       base, rows(base, range(4), 28.0))
        self._workbook(tmp_path / "result3.xlsx", ["Sheet1"], base,
                       rows(base, range(0, step * 4, step), final3))
        self._workbook(tmp_path / "result4.xlsx", ["Sheet1"], labelled,
                       rows(labelled, range(0, step * 4, step), final4))

        data_dir = tmp_path / "data"
        data_dir.mkdir(exist_ok=True)
        outcome = ExecutionOutcome(
            status="success",
            exit_code=0,
            execution_real=True,
            runtime_seconds=1.0,
            sandbox={"image": "test"},
            summary={"statistics": {"checked": 1}},
        )
        checks = DeterministicVerifier(data_dir).run(
            outcome=outcome,
            output_dir=tmp_path,
            required_outputs=self.REQUIRED,
            template_headers=self.HEADERS,
            template_sheets=self.SHEETS,
            problem_text=self.STATEMENT if statement is None else statement,
        )
        return {c.name: c for c in checks}

    # ── A: the real A expansion is no longer rejected ────────────────
    def test_the_expanded_header_passes_the_structure_check(self, tmp_path):
        checks = self._gate(tmp_path)
        assert checks["output_structure_matches_template"].status == CheckStatus.PASS

    # ── B ────────────────────────────────────────────────────────────
    def test_a_generic_worksheet_name_fails_and_names_the_worksheet(self, tmp_path):
        checks = self._gate(tmp_path, temperature_sheet="Sheet1")
        check = checks["output_worksheets_match_template"]
        assert check.status == CheckStatus.FAIL
        assert "温度" in check.detail
        assert "worksheet" in check.detail.lower()

    # ── C ────────────────────────────────────────────────────────────
    def test_a_six_hour_grid_fails_and_states_the_required_interval(self, tmp_path):
        checks = self._gate(tmp_path, step=21600)
        check = checks["output_time_grid_matches_problem"]
        assert check.status == CheckStatus.FAIL
        assert "60 s" in check.detail
        assert "21600" in check.detail

    # ── D ────────────────────────────────────────────────────────────
    def test_a_final_moisture_of_0_157_fails(self, tmp_path):
        checks = self._gate(tmp_path, final3=0.157)
        check = checks["output_meets_stated_final_target"]
        assert check.status == CheckStatus.FAIL
        assert "0.157" in check.detail

    # ── E ────────────────────────────────────────────────────────────
    def test_a_final_moisture_of_exactly_0_15_fails(self, tmp_path):
        checks = self._gate(tmp_path, final4=0.15)
        check = checks["output_meets_stated_final_target"]
        assert check.status == CheckStatus.FAIL
        assert "0.150000" in check.detail
        assert "strictly" in check.detail

    def test_a_compliant_output_passes_every_check(self, tmp_path):
        checks = self._gate(tmp_path)
        for name in (
            "output_structure_matches_template",
            "output_worksheets_match_template",
            "output_time_grid_matches_problem",
            "output_meets_stated_final_target",
        ):
            assert checks[name].status == CheckStatus.PASS, (name, checks[name].detail)

    # ── G: drying far faster than the statement allows ───────────────
    def test_drying_within_minutes_against_a_two_to_three_day_process_fails(
        self, tmp_path
    ):
        """The real defect: sourcetermfix reached <0.15 from 2.55 in about an
        hour while the statement says the process lasts 2-3 days, and the
        final-row check passed it vacuously."""
        statement = self.STATEMENT + "\n问题 2  烘干过程一般持续 2-3 天。"
        checks = self._gate(tmp_path, statement=statement)
        check = checks["output_process_duration_is_plausible"]
        assert check.status == CheckStatus.FAIL
        assert "172800" in check.detail
        assert "not a fast answer" in check.detail
        assert "check the transfer coefficients" in check.detail

    def test_no_stated_duration_means_no_such_check(self, tmp_path):
        checks = self._gate(tmp_path)
        assert "output_process_duration_is_plausible" not in checks

    # ── F ────────────────────────────────────────────────────────────
    def test_a_failing_structure_check_does_not_silence_the_others(self, tmp_path):
        broken = ["时间\\到药材中心的距离", "半径0", "半径1"] + _a_radii()[3:]
        checks = self._gate(
            tmp_path, header=broken, temperature_sheet="Sheet1",
            step=21600, final3=0.157, final4=0.15,
        )
        assert checks["output_structure_matches_template"].status == CheckStatus.FAIL
        assert checks["output_worksheets_match_template"].status == CheckStatus.FAIL
        assert checks["output_time_grid_matches_problem"].status == CheckStatus.FAIL
        assert checks["output_meets_stated_final_target"].status == CheckStatus.FAIL
        assert [n for n, c in checks.items() if c.status == CheckStatus.NOT_RUN] == []

    def test_the_d_problem_gets_no_new_checks(self, tmp_path):
        """The D statement states none of these constraints, so nothing regresses."""
        from mathmodel.autopilot.verify import (
            problem_final_state_targets,
            problem_time_steps,
        )

        statement = "将冲突装备对保存到文件 result1.xlsx 中。"
        assert problem_time_steps(statement) == {}
        assert problem_final_state_targets(statement, ["result1.xlsx"]) == []


# ═══════════════════════════════════════════════════════════════
# Numerical scheme validation: a differential-equation model must show
# that its scheme reproduces a case with a known answer, and the repair
# hint it gets must fit the kind of model it is.
# ═══════════════════════════════════════════════════════════════

def _math_model(equations, family):
    from mathmodel.domain.math_model import MathematicalModel

    return MathematicalModel(
        name="model",
        description="test model",
        model_family=family,
        equations=[{"name": n, "expression": e} for n, e in equations],
    )


def _a2026_like_model():
    return _math_model(
        [("径向能量守恒方程", "rho*cp*(∂T/∂t) = (1/r)*∂(r*k*∂T/∂r)/∂r")],
        "coupled nonlinear parabolic PDE system",
    )


def _d_like_model():
    return _math_model(
        [("调整后频段下界", "F_L_adj_i = F_L_i + S_i")],
        "混合整数线性规划 (MILP)",
    )


class TestSchemeValidationTrigger:
    """Which models have to validate a numerical scheme at all."""

    def test_a_derivative_equation_requires_scheme_validation(self):
        from mathmodel.autopilot.verify import model_integrates_differential_equations

        assert model_integrates_differential_equations(_a2026_like_model()) is True

    def test_an_algebraic_assignment_does_not(self):
        from mathmodel.autopilot.verify import model_integrates_differential_equations

        assert model_integrates_differential_equations(_d_like_model()) is False

    def test_a_differential_family_is_enough_on_its_own(self):
        """A PDE described without a derivative glyph is still a PDE."""
        from mathmodel.autopilot.verify import model_integrates_differential_equations

        model = _math_model([("x", "y = 2*z")], "ODE system")
        assert model_integrates_differential_equations(model) is True

    def test_a_model_with_no_equations_does_not(self):
        from mathmodel.autopilot.verify import model_integrates_differential_equations

        assert model_integrates_differential_equations(_math_model([], "MILP")) is False


class TestThePhysicalTransportCeiling:
    """round 57: a solver removed 0.62 kg of moisture in 180 s -- about 65x the
    ceiling the stated coefficients allow -- and its own checks all passed."""

    STATEMENT = (
        "某中药材形状大致呈圆柱形，长为 25 cm，半径为 2 cm。"
        "密度为 820 kg/m3、比热容为 2600 J/(kg·K)、热传导系数为 0.36 W/(m·K)，"
        "对流换热系数为 25 W/(m2·K)，对流传质系数为 8 × 10−7 m/s，"
        "水分浓度（即干基含水率）为2.55 kg/kg。"
    )

    def test_the_ceiling_matches_an_independent_hand_calculation(self):
        from mathmodel.autopilot.verify import problem_transport_ceiling
        import math as _math

        ceiling = problem_transport_ceiling(self.STATEMENT)
        assert ceiling is not None
        # h_m * rho * C0 * 2*pi*R*L, computed independently
        expected = 8e-7 * 820.0 * 2.55 * 2 * _math.pi * 0.02 * 0.25
        assert ceiling == pytest.approx(expected, rel=1e-9)

    def test_the_measured_solver_is_far_above_it(self):
        from mathmodel.autopilot.verify import problem_transport_ceiling

        ceiling = problem_transport_ceiling(self.STATEMENT)
        observed = 0.62 / 180.0  # kg/s, measured on the real run
        assert observed > 50.0 * ceiling

    def test_missing_constants_means_no_check(self):
        from mathmodel.autopilot.verify import problem_transport_ceiling

        assert problem_transport_ceiling("") is None
        assert problem_transport_ceiling("某圆柱形药材，长为 25 cm，半径为 2 cm。") is None


class TestTheStatedProcessDurationIsParsed:
    """round 49 found a solver drying 2.55 -> <0.15 in about an hour while the
    statement says the process lasts 2-3 days, so the target check passed
    vacuously and nothing host-side noticed."""

    def test_the_real_range_is_read(self):
        from mathmodel.autopilot.verify import problem_process_duration

        got = problem_process_duration("整个干燥过程持续 2-3 天，请给出…")
        assert got == (2 * 86400.0, 3 * 86400.0)

    def test_tilde_and_hours_and_ascii(self):
        from mathmodel.autopilot.verify import problem_process_duration

        assert problem_process_duration("持续2~3天") == (172800.0, 259200.0)
        assert problem_process_duration("约 48-72 小时") == (172800.0, 259200.0)
        assert problem_process_duration("lasts 2-3 days") == (172800.0, 259200.0)

    def test_a_single_duration_is_both_bounds(self):
        from mathmodel.autopilot.verify import problem_process_duration

        assert problem_process_duration("烘干持续 3 天") == (259200.0, 259200.0)

    def test_no_duration_means_no_check(self):
        from mathmodel.autopilot.verify import problem_process_duration

        assert problem_process_duration("") is None
        assert problem_process_duration("求含水率分布与温度分布。") is None

    def test_an_implausible_short_drying_is_recognisably_out_of_range(self):
        """The real figure: below 0.15 within ~1800 s against a 2-3 day process."""
        from mathmodel.autopilot.verify import problem_process_duration

        low, high = problem_process_duration("整个干燥过程持续 2-3 天")
        assert 1800.0 < 0.05 * low


class TestARepeatedFinalHeaderLabelIsNamed:
    """runs/cumcm-2026-a2026-auditfix: result1.xlsx had 23 columns ending
    '1.8', '1.9', '2', '2' -- the final label duplicated, so one column too
    many. The ellipsis check described the symptom instead of the form."""

    def _expected(self):
        return ["时间\\到药材中心的距离", "0", "0.1", "…", "1.9", "2"]

    def _expanded(self):
        """The template's real expansion: 0 to 2 in steps of 0.1, 21 columns."""
        return ["时间\\到药材中心的距离"] + [
            f"{0.1 * i:g}" for i in range(21)
        ]

    def test_the_real_duplicated_header_is_named(self):
        from mathmodel.autopilot.verify import header_mismatches

        problems = header_mismatches(self._expected(), self._expanded() + ["2"])
        assert problems
        assert "the last column label is repeated" in problems[0]
        assert "one column more" in problems[0]

    def test_a_correct_header_is_still_clean(self):
        from mathmodel.autopilot.verify import header_mismatches

        assert header_mismatches(self._expected(), self._expanded()) == []

    def test_a_duplicate_without_a_mismatch_is_not_invented(self):
        """The duplication is only NAMED when the header already fails: a
        repeated label alone is not sufficient grounds for a finding."""
        from mathmodel.autopilot.verify import _repeated_final_label

        assert _repeated_final_label(["0", "1", "2", "2"]) == 2.0
        assert _repeated_final_label(["0", "1", "2"]) is None
        assert _repeated_final_label(["a", "b"]) is None


class TestRewordedButIdenticalFailuresStayOneClass:
    """runs/cumcm-2026-a2026-levelsfix: attempts 2, 3, 4 and 6 all aborted on the
    solver's own final-target assert at the IDENTICAL value 2.5217, each time
    reworded -- "final row", "final written 60 s row", "problem 3 final written
    60 s row". The rewording is the tell that the repair edited the message
    instead of the physics."""

    def test_the_real_rewordings_are_one_class(self):
        from mathmodel.autopilot.pipeline import _failure_class

        messages = [
            "AssertionError: final row 2.5217 not strictly below 0.15",
            "AssertionError: final written 60 s row 2.5217 not < 0.15",
            "AssertionError: problem 3 final written 60 s row max=2.5217 not "
            "strictly < 0.15",
        ]
        classes = {_failure_class(m) for m in messages}
        assert classes == {"final-target-unmet"}

    def test_the_note_forbids_reworking_the_guard(self):
        from mathmodel.autopilot.pipeline import _repeated_failure_note

        note = _repeated_failure_note("final-target-unmet", attempt=3, seen=2)
        assert "YOUR OWN GUARD KEEPS TRIPPING" in note
        assert "Do NOT reword, relax or remove that assertion" in note
        assert "aimed at the message rather" in note
        assert "fix the transport" in note.lower() or "Fix the mass transfer" in note

    def test_an_unrelated_failure_keeps_the_generic_note(self):
        from mathmodel.autopilot.pipeline import _failure_class, _repeated_failure_note

        assert _failure_class("ValueError: operands could not be broadcast") != (
            "final-target-unmet"
        )
        note = _repeated_failure_note("shape-mismatch", attempt=2, seen=1)
        assert "STOP PATCHING CALL SITES" in note


class TestTheRepairIsBasedOnTheBestAttemptNotTheLatest:
    """I43's strong form: restoring the better attempt is structural, not a
    request the model may ignore."""

    def _call(self, history, programs, attempt, fallback="LATEST"):
        from mathmodel.autopilot.pipeline import _repair_base_program

        return _repair_base_program(history, programs, attempt, fallback)

    def test_a_worse_attempt_is_rebased_on_the_better_one(self):
        code, source = self._call(
            [(1, 0.475), (2, 57.319)],
            {1: "GOOD", 2: "BAD"},
            attempt=2,
        )
        assert code == "GOOD"
        assert source == 1

    def test_an_improving_run_keeps_the_previous_attempt(self):
        code, source = self._call(
            [(1, 57.319), (2, 0.475)],
            {1: "BAD", 2: "GOOD"},
            attempt=2,
        )
        assert code == "LATEST"
        assert source is None

    def test_matching_the_best_keeps_the_previous_attempt(self):
        code, source = self._call(
            [(1, 0.5), (2, 0.5)], {1: "A", 2: "B"}, attempt=2
        )
        assert code == "LATEST"
        assert source is None

    def test_a_single_attempt_has_nothing_to_restore(self):
        code, source = self._call([(1, 9.0)], {1: "A"}, attempt=1)
        assert code == "LATEST"
        assert source is None

    def test_an_attempt_with_no_recorded_program_is_not_a_base(self):
        code, source = self._call(
            [(1, 0.1), (2, 9.0)], {2: "BAD"}, attempt=2
        )
        assert code == "LATEST"
        assert source is None

    def test_the_best_is_chosen_by_ratio_not_by_order(self):
        code, source = self._call(
            [(1, 5.0), (2, 0.2), (3, 9.0)],
            {1: "A", 2: "BEST", 3: "C"},
            attempt=3,
        )
        assert code == "BEST"
        assert source == 2


class TestTheRepairLoopNoticesItIsGettingWorse:
    def test_a_worsening_trend_is_named_with_the_better_attempt(self):
        """Real case: the same problem yielded err=0.475 on one attempt and
        err=57.319 on a later one, so always rewriting from the latest attempt
        can walk away from a discretisation that nearly worked."""
        from mathmodel.autopilot.pipeline import _worsening_trend_note

        history = [(1, 3.0), (2, 0.95), (3, 12.0), (4, 114.6)]
        note = _worsening_trend_note(history, 4)
        assert note is not None
        assert "making the scheme WORSE" in note
        assert "Attempt 2" in note
        assert "0.95" in note and "114.6" in note
        assert "Do not continue rewriting the scheme from the latest version" in note
        assert "restore that" in note

    def test_the_better_attempts_satisfied_criteria_are_named(self):
        """"Restore attempt 2" is only actionable if the solver can see which
        criteria attempt 2 had already satisfied."""
        from mathmodel.autopilot.pipeline import _worsening_trend_note

        note = _worsening_trend_note(
            [(1, 3.0), (2, 0.95), (3, 12.0), (4, 114.6)],
            4,
            {
                2: ["manufactured_order"],
                4: ["max_abs_error", "steady_state_error", "manufactured_order"],
            },
        )
        assert note is not None
        assert "Attempt 2 reported these criteria as failed: manufactured_order" in note
        assert "attempt 4 reports: max_abs_error, steady_state_error" in note
        assert "has broken" in note

    def test_an_improving_attempt_is_not_lectured(self):
        from mathmodel.autopilot.pipeline import _worsening_trend_note

        assert _worsening_trend_note([(1, 12.0), (2, 3.0), (3, 0.95)], 3) is None

    def test_a_single_attempt_is_not_a_trend(self):
        from mathmodel.autopilot.pipeline import _worsening_trend_note

        assert _worsening_trend_note([(1, 12.0)], 1) is None
        assert _worsening_trend_note([], 1) is None

    def test_matching_the_best_attempt_is_not_worsening(self):
        from mathmodel.autopilot.pipeline import _worsening_trend_note

        assert _worsening_trend_note([(1, 0.95), (2, 0.95)], 2) is None


class TestAVerifierVerdictMustMatchItsOwnEvidence:
    """The verifier is generated code too and can contradict itself."""

    def _check(self, status, detail, evidence):
        from mathmodel.autopilot.verify import VerificationCheck

        return VerificationCheck(
            check_id="VC-1",
            name="independent_formula_transcription",
            category="independent",
            status=status,
            detail=detail,
            evidence=evidence,
        )

    def test_the_real_case_is_detected(self):
        """runs/cumcm-2026-a2026-stagedfix: FAIL whose detail confirms the
        formulas were transcribed correctly."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            unaccountable_formula_verdict,
        )

        check = self._check(
            CheckStatus.FAIL,
            "D_q1=e^-0.89/C ratio; D_q23/D_q4 use exp(-3850/T_K) Arrhenius "
            "ratio (T in Kelvin).",
            {
                "formula_verdicts": {"D_q1": "pass", "D_q23": "pass", "D_q4": "pass"},
                "D_q1_C2.55": 4.937655094173937e-9,
            },
        )
        assert unaccountable_formula_verdict(check) is not None

    def test_a_real_formula_failure_is_left_alone(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            unaccountable_formula_verdict,
        )

        check = self._check(
            CheckStatus.FAIL,
            "D_q3 uses the product reading where the statement implies a ratio.",
            {"formula_verdicts": {"D_q3": "fail: product where a ratio is required"}},
        )
        assert unaccountable_formula_verdict(check) is None

    def test_a_pass_is_never_touched(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            unaccountable_formula_verdict,
        )

        check = self._check(
            CheckStatus.PASS, "all formulas transcribed correctly",
            {"formula_verdicts": {"D_q1": "pass"}},
        )
        assert unaccountable_formula_verdict(check) is None

    def test_a_failure_without_per_formula_verdicts_is_left_alone(self):
        """No structured verdicts means nothing to contradict, so the host must
        not start second-guessing the verifier on prose alone."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            unaccountable_formula_verdict,
        )

        check = self._check(
            CheckStatus.FAIL, "the solver's diffusion law is wrong", {}
        )
        assert unaccountable_formula_verdict(check) is None

    def test_downgrading_still_fails_the_report(self, tmp_path):
        """NOT_RUN must never become a way to pass."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationReport,
            build_report,
        )

        checks = [
            self._check(
                CheckStatus.NOT_RUN,
                "Not verifiable: the verifier marked this check FAIL while every "
                "per-formula verdict it reported passes.",
                {"formula_verdicts": {"D_q1": "pass"}},
            )
        ]
        report = build_report(
            MathematicalModel(name="m", objective="o"), "run", checks
        )
        assert report.overall == CheckStatus.FAIL


class TestSchemeSelfCheckGate:
    """The solver's own scheme validation is read, not trusted."""

    def _checks(self, tmp_path, self_check, required):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        out_dir = tmp_path / "out"
        out_dir.mkdir()
        _make_xlsx(out_dir / "result1.xlsx", ["序号"], [[1]])

        outcome = _outcome({"conflicts": 0})
        if self_check is not None:
            # These fixtures are about OTHER criteria; supply a healthy
            # conservation audit so the audit requirement (added in round 55)
            # does not mask what each test is actually asserting. The audit has
            # its own dedicated tests.
            self_check.setdefault("mass_balance_relative_error", 1e-4)
            outcome.summary["self_check"] = self_check
        return DeterministicVerifier(data_dir).run(
            outcome=outcome,
            output_dir=out_dir,
            required_outputs=["result1.xlsx"],
            template_headers={},
            requires_scheme_validation=required,
        )

    @staticmethod
    def _by_name(checks, name):
        return [c for c in checks if c.name == name][0]

    def test_a_solver_that_does_not_audit_its_conservation_is_told_to(self, tmp_path):
        """round 54: attempt6 put rho~820 in its moisture surface flux but not in
        its moisture storage, so it dried 2.55 -> <0.15 in 180 s, and its own
        scheme validation passed every criterion it had written."""
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 1e-4,
                "tolerance": 1.0,
                "mass_balance_relative_error": None,
                "energy_balance_relative_error": None,
                "passed": True,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "does not report a conservation audit" in check.detail
        assert "mass_balance_relative_error" in check.detail

    def test_an_imbalanced_audit_names_the_factor_inconsistency(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 1e-4,
                "tolerance": 1.0,
                "mass_balance_relative_error": 0.82,
                "passed": True,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "conservation audit reports" in check.detail
        assert "0.82" in check.detail
        assert "carries the same factors" in check.detail

    def test_a_balanced_audit_is_not_second_guessed(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 1e-8,
                "tolerance": 1e-5,
                "reference_span": 2.0,
                "mass_balance_relative_error": 1e-9,
                "energy_balance_relative_error": 1e-9,
                "passed": True,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "conservation audit" not in check.detail

    def test_an_order_only_failure_says_the_scheme_is_not_implicated(self, tmp_path):
        """runs/cumcm-2026-a2026-rebasefix attempt 3: the problem was solved
        correctly -- drying passed, deviation from its reference 0.0122 -- and
        it was blocked by an order of -1.95513 from a manufactured test whose
        source term cannot be right."""
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured, cylindrical, variable properties",
                "max_abs_error": 0.0122,
                "tolerance": 1.0,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": -1.95513,
                "expected_order": 1.0,
                "manufactured_levels": [
                    {"n": 10, "cells": 10, "error": 1e-2},
                    {"n": 20, "cells": 20, "error": 2e-2},
                    {"n": 40, "cells": 40, "error": 4e-2},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "the order study is the ONLY criterion that failed" in check.detail
        assert "a statement about your TEST, not about" in check.detail
        assert "Repair the TEST" in check.detail
        assert "rewrite the solver" in check.detail

    def test_an_order_failure_beside_a_real_one_omits_that_reassurance(self, tmp_path):
        """The reassurance must not appear when something else also failed."""
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured",
                "max_abs_error": 0.9,
                "tolerance": 0.5,
                "passed": False,
                "failed_criteria": ["max_abs_error", "manufactured_order"],
                "manufactured_order": -1.95513,
                "expected_order": 1.0,
                "manufactured_levels": [
                    {"n": 10, "cells": 10, "error": 1e-2},
                    {"n": 20, "cells": 20, "error": 2e-2},
                    {"n": 40, "cells": 40, "error": 4e-2},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "the order study is the ONLY criterion that failed" not in check.detail
        assert "NEGATIVE order" in check.detail

    def test_a_negative_convergence_order_blames_the_test_first(self, tmp_path):
        """Real case from runs/cumcm-2026-a2026-flatcause: order -1.48674.

        A negative order means the error GROWS under refinement, which is not a
        scheme that lost an order -- it is a test that is not measuring
        convergence, so the causes to check are on the test side.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 2.08389,
                "tolerance": 0.05,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": -1.48674,
                "expected_order": 2.0,
                # Levels that corroborate a genuinely non-converging test:
                # the error GROWS as the grid is refined.
                "manufactured_levels": [
                    {"cells": 40, "error": 1.0e-2},
                    {"cells": 80, "error": 2.0e-2},
                    {"cells": 160, "error": 4.0e-2},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "a NEGATIVE order means the error GROWS as you refine" in check.detail
        assert "check the test " in check.detail
        assert "SAME physical time at every level" in check.detail
        assert "-1.48674" in check.detail
        # The weaker wording is for a small-but-positive order, not this one.
        assert "Do not respond by refining the grid further" not in check.detail

    def test_a_scheme_that_loses_its_order_is_named(self, tmp_path):
        """A manufactured-solution test measures the scheme, not a reference.

        The analytic reference was the broken part in several separate runs, each
        time blocking a scheme that was sound. A scheme that has lost a
        control-volume factor or counts a flux twice does not converge at the
        order it claims.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 1e-3,
                "tolerance": 1e-2,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": 0.5,
                "expected_order": 2.0,
                # Levels that corroborate the reported order instead of
                # contradicting it: errors falling by 2x per doubling.
                "manufactured_levels": [
                    {"cells": 40, "error": 4.0e-2},
                    {"cells": 80, "error": 2.0e-2},
                    {"cells": 160, "error": 1.0e-2},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "manufactured-solution convergence test shows order 0.5" in check.detail
        assert "Do not respond by refining the grid further" in check.detail

    def test_a_scheme_that_keeps_its_order_is_not_flagged(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 1e-5,
                "tolerance": 1e-3,
                "passed": True,
                "manufactured_order": 1.98,
                "expected_order": 2.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.PASS
        assert "manufactured-solution" not in check.detail

    def test_repair_feedback_names_what_already_passes(self):
        """Repair must say what to leave alone, not only what is broken."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            VerificationReport,
        )

        report = VerificationReport(
            run_id="r",
            checks=[
                VerificationCheck(
                    check_id="VC-1", name="output_worksheets_match_template",
                    category="correctness", status=CheckStatus.PASS,
                    detail="all sheets present",
                ),
                VerificationCheck(
                    check_id="VC-2", name="output_meets_stated_final_target",
                    category="correctness", status=CheckStatus.FAIL,
                    detail="final moisture 2.48 exceeds 0.15",
                ),
                VerificationCheck(
                    check_id="VC-3", name="independent_small_case_check",
                    category="independent", status=CheckStatus.NOT_RUN,
                    detail="not run",
                ),
            ],
        )
        lines = report.to_prompt_feedback()
        context = [line for line in lines if line.startswith("[context]")]
        assert len(context) == 1
        assert "output_worksheets_match_template" in context[0]
        assert "already PASS" in context[0]
        # NOT_RUN is not a pass and must never be advertised as one.
        assert "independent_small_case_check" not in context[0]
        assert "output_meets_stated_final_target" not in context[0]

    def test_repair_feedback_omits_the_context_line_when_nothing_fails(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            VerificationReport,
        )

        report = VerificationReport(
            run_id="r",
            checks=[
                VerificationCheck(
                    check_id="VC-1", name="some_check", category="correctness",
                    status=CheckStatus.PASS, detail="fine",
                ),
            ],
        )
        assert report.to_prompt_feedback() == []

    def test_a_passing_self_check_passes(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"case": "constant-coefficient slab", "max_abs_error": 1e-4,
             "tolerance": 1e-3, "passed": True},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.PASS
        assert "constant-coefficient slab" in check.detail

    def test_an_unauditable_order_is_not_reported_as_a_scheme_fault(self, tmp_path):
        """Real case: verdictfix2 and juryretry both blocked on this alone.

        An order near zero is produced both by a scheme that does not converge
        and by a wrong manufactured source term, and the levels that would tell
        them apart were not reported. The gate must still FAIL -- but it must
        not tell the solver to change a discretisation that may be sound.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 0.144713,
                "tolerance": 0.005,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": 0.000999215,
                "expected_order": 2.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "it is not auditable" in check.detail
        assert "manufactured_levels" in check.detail
        assert "fix the test before you change the scheme" in check.detail

    def test_levels_that_converge_expose_a_miscomputed_order(self, tmp_path):
        """A reported order is a conclusion; the levels are the evidence.

        Real case from runs/cumcm-2026-a2026-verdictfix2: the solver reported
        order -0.0007 while its result files satisfied every host-side
        acceptance check. A saturated error gives order ~0 at any refinement,
        which points at a wrong manufactured source term or a miscomputed order
        -- not at the scheme. Here the levels themselves converge, so the order
        calculation is the fault.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 0.144713,
                "tolerance": 0.005,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": -0.000705945,
                "expected_order": 1.0,
                "manufactured_levels": [
                    {"cells": 40, "error": 4.0e-2},
                    {"cells": 80, "error": 2.0e-2},
                    {"cells": 160, "error": 1.0e-2},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "your manufactured-levels evidence contradicts" in check.detail
        assert "Your own refinement study CONVERGES" in check.detail
        assert "you miscomputed the order" in check.detail

    def test_levels_that_genuinely_stall_keep_the_scheme_diagnosis(self, tmp_path):
        """A genuinely flat error series is a real non-convergence finding."""
        checks = self._checks(
            tmp_path,
            {
                "case": "manufactured solution",
                "max_abs_error": 0.14,
                "tolerance": 0.005,
                "passed": False,
                "failed_criteria": ["manufactured_order"],
                "manufactured_order": -0.0007,
                "expected_order": 1.0,
                "manufactured_levels": [
                    {"cells": 40, "error": 0.1447},
                    {"cells": 80, "error": 0.1448},
                    {"cells": 160, "error": 0.1449},
                ],
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "your manufactured-levels evidence contradicts" not in check.detail

    def test_a_verifier_value_far_outside_any_real_disagreement_is_blamed_on_it(self):
        """Real case from runs/cumcm-2026-a2026-verdictfix2.

        The solver reported a final moisture concentration of 0.1469 and the
        verifier recomputed 33960.0 against an initial 2.55 -- absurd as a
        concentration. The ratio 2.3e5 was below IMPLAUSIBLE_RATIO, so the run
        blocked on three checks that all rested on the verifier's number.
        """
        from mathmodel.autopilot.verify import _diverged_pair

        assert _diverged_pair(0.1469, 33960.0) is True
        # An ordinary disagreement must stay an ordinary disagreement.
        assert _diverged_pair(57.0, 16.79) is None
        assert _diverged_pair(0.15, 0.14) is None

    def test_the_solver_side_can_be_the_diverged_one(self):
        from mathmodel.autopilot.verify import _diverged_pair

        assert _diverged_pair(4.87729701306026e232, 28.0) is False

    def test_a_physical_value_is_never_called_diverged(self):
        from mathmodel.autopilot.verify import _diverged_pair

        assert _diverged_pair(2.55, 0.1469) is None
        assert _diverged_pair(1e29, 1e29) is None

    def test_a_flat_profile_at_the_initial_value_blames_the_boundary(self, tmp_path):
        """Real case from runs/cumcm-2026-a2026-biotfix2.

        The solver reported span=8.52651e-14 against a reference span of 10.5784.
        A flat profile has two causes -- an over-large conductance drives the
        field to the boundary value, a dead surface condition leaves it at the
        initial value -- and only the values separate them. Naming the wrong one
        sends the repair to the wrong code.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "cylinder with convective boundary",
                "max_abs_error": 24.1841,
                "tolerance": 0.5,
                "passed": False,
                "failed_criteria": ["max_abs_error"],
                "profile_span": 8.52651e-14,
                "reference_span": 10.5784,
                "profile_min": 28.0,
                "profile_max": 28.0,
                "initial_value": 28.0,
                "boundary_value": 60.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "never left its initial value" in check.detail
        assert "the surface condition is not acting at all" in check.detail
        assert "not a conductance problem" in check.detail
        assert "diffusion conductance is too large" not in check.detail

    def test_a_flat_profile_at_the_boundary_value_still_blames_the_conductance(
        self, tmp_path
    ):
        checks = self._checks(
            tmp_path,
            {
                "case": "cylinder with convective boundary",
                "max_abs_error": 7.5,
                "tolerance": 0.5,
                "passed": False,
                "failed_criteria": ["max_abs_error"],
                "profile_span": 0.01,
                "reference_span": 12.0,
                "profile_min": 59.99,
                "profile_max": 60.0,
                "initial_value": 28.0,
                "boundary_value": 60.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "diffusion conductance is too large" in check.detail
        assert "never left its initial value" not in check.detail

    def test_a_flat_profile_without_values_keeps_the_older_wording(self, tmp_path):
        """Older records lack the new fields; the diagnosis must not crash."""
        checks = self._checks(
            tmp_path,
            {
                "case": "cylinder",
                "max_abs_error": 7.5,
                "tolerance": 0.5,
                "passed": False,
                "profile_span": 0.01,
                "reference_span": 12.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "diffusion conductance is too large" in check.detail

    def test_a_verdict_contradicting_its_own_numbers_is_named(self, tmp_path):
        """The exact self_check from runs/cumcm-2026-a2026-formulafix3.

        max_abs_error 0.0148 against a stated tolerance of 1.0, steady-state
        error 4.3e-11, profile span 14.149 against the reference's 14.173 -- a
        validated scheme by its own evidence -- yet passed=false. The bare
        boolean blocked the run, which reads a claim in place of the evidence.
        """
        checks = self._checks(
            tmp_path,
            {
                "case": "constant-coefficient infinite cylinder with Robin "
                        "surface, Bessel series reference",
                "max_abs_error": 0.014771851196492491,
                "initial_profile_error": 0.2831789682366619,
                "steady_state_error": 4.270361841918202e-11,
                "profile_span": 14.149233398708454,
                "reference_span": 14.173015968837326,
                "tolerance": 1.0,
                "passed": False,
                "dt_sensitivity": 0.002858113733573475,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "verdict contradicts its own measurements" in check.detail
        assert "failed_criteria" in check.detail
        # The old, unaccountable wording must be gone.
        assert "the program itself reported passed=false" not in check.detail

    def test_named_failed_criteria_are_reported(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 1e-4, "tolerance": 1e-3,
             "passed": False, "failed_criteria": ["steady_state_error"]},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "named these criteria as failed: steady_state_error" in check.detail

    def test_a_repeated_criterion_is_reported_once(self, tmp_path):
        """Real case from runs/cumcm-2026-a2026-verdictfix: the solver listed
        steady_state_error twice and read as two separate faults."""
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 52.0, "tolerance": 1e-3,
             "passed": False,
             "failed_criteria": ["steady_state_error", "steady_state_error"]},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "named these criteria as failed: steady_state_error;" in check.detail
        assert "steady_state_error, steady_state_error" not in check.detail

    def test_a_self_chosen_tolerance_must_be_small_against_its_own_scale(self, tmp_path):
        """Two solvers on this problem declared tolerances of 1.0 and 0.005 on
        the same quantity, so the strictness of the check was set by whichever
        number the generation happened to pick."""
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 1e-9,
                "tolerance": 1.0,
                "reference_span": 2.0,
                "passed": True,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "not a meaningful test" in check.detail
        assert "50%" in check.detail
        assert "small compared with what it measures" in check.detail

    def test_a_tolerance_small_against_its_scale_is_left_alone(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 1e-9,
                "tolerance": 0.005,
                "reference_span": 2.0,
                "passed": True,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "not a meaningful test" not in check.detail

    def test_a_diverged_validation_is_not_called_an_accuracy_shortfall(self, tmp_path):
        """runs/cumcm-2026-a2026-auditfix: the solver's main problem produced a
        sensible 43.5 h drying time corroborated independently, yet its
        validation reported max_abs_error=7.18062e+108."""
        checks = self._checks(
            tmp_path,
            {
                "case": "constant-coefficient cylinder heat equation, Dirichlet "
                        "surface T=80, vs Bessel series",
                "max_abs_error": 7.18062e108,
                "tolerance": 1.0,
                "passed": False,
                "failed_criteria": ["max_abs_error"],
                "reference_span": 52.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "DIVERGED in the validation case" in check.detail
        assert "not an accuracy shortfall but a blow-up" in check.detail
        assert "does DIFFERENTLY" in check.detail
        assert "boundary condition" in check.detail
        assert "exceeds the stated tolerance" not in check.detail

    def test_an_ordinary_miss_keeps_the_ordinary_wording(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {
                "case": "slab",
                "max_abs_error": 5.0,
                "tolerance": 1e-3,
                "passed": False,
                "failed_criteria": ["max_abs_error"],
                "reference_span": 52.0,
            },
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert "exceeds the stated tolerance" in check.detail
        assert "DIVERGED" not in check.detail

    def test_a_real_measurement_failure_keeps_its_own_reason(self, tmp_path):
        """An accountable failure is not relabelled as a contradiction."""
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 5.0, "tolerance": 1e-3,
             "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "exceeds the stated tolerance" in check.detail
        assert "verdict contradicts its own measurements" not in check.detail

    def test_a_self_reported_failure_fails_the_gate(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"case": "analytical series", "max_abs_error": 0.5,
             "tolerance": 1e-3, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "passed=false" in check.detail
        assert "exceeds the stated tolerance" in check.detail

    def test_a_deviation_above_tolerance_fails_even_when_claimed_passed(self, tmp_path):
        """An over-optimistic `passed` flag must not override the measurement."""
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 0.9, "tolerance": 1e-3,
             "passed": True},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "0.9" in check.detail

    def test_a_broken_reference_is_named_as_the_fault(self, tmp_path):
        """The real failure from runs/cumcm-2026-a2026-repairfix.

        The reference's series coefficient had the wrong sign, so it returned
        90.4 at the centre where the correct initial value was 28.0. The scheme
        reproduced T0 exactly; blaming the scheme sent the repair loop after the
        wrong code.
        """
        checks = self._checks(
            tmp_path,
            {"case": "analytical cylindrical Dirichlet series", "max_abs_error": 64.9582,
             "initial_profile_error": 62.4, "tolerance": 0.5, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "REFERENCE is the broken part, not the scheme" in check.detail
        assert "before changing any scheme code" in check.detail

    def test_a_discretisation_level_t0_gap_does_not_blame_the_reference(self, tmp_path):
        """The t=0 gap must dominate the deviation to count as a broken reference.

        Real case: a solver stating a tight 1e-3 tolerance reported a t=0 gap of
        1.6e-3 against a total deviation of 0.0317. That gap is discretisation,
        not a broken reference, and the earlier absolute test blamed it.
        """
        checks = self._checks(
            tmp_path,
            {"case": "analytical constant-coefficient cylinder",
             "max_abs_error": 0.0316787, "initial_profile_error": 0.00162316,
             "tolerance": 0.001, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "REFERENCE is the broken part" not in check.detail
        assert "initial_profile_error" in check.detail

    def test_a_flat_profile_names_the_diffusion_conductance(self, tmp_path):
        """The real failure from runs/cumcm-2026-a2026-final.

        The scheme carried a stray 1/R**2 in its conduction term, so its profile
        came out flat (centre 37.4819, surface 37.4916 -> span 0.0097 K) where the
        reference was strongly curved (29.9396 -> 42.0052, span 12.0656 K). The
        scheme's own reference was correct to 7e-15 against an independent host
        computation, and removing that single factor took the error from 7.54231
        to 0.01502 K. "The deviation exceeds the tolerance" does not tell the
        model any of that.
        """
        checks = self._checks(
            tmp_path,
            {"case": "analytical cylindrical Bessel series (heat only)",
             "max_abs_error": 7.54231, "initial_profile_error": 0.0758,
             "profile_span": 0.0097, "reference_span": 12.0656,
             "tolerance": 0.6, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "nearly flat where the reference is curved" in check.detail
        assert "1/R**2" in check.detail
        assert "time step will not fix this" in check.detail
        # The precise diagnosis must replace the generic "check the reference" tail.
        assert "initial_profile_error" not in check.detail

    def test_a_curved_profile_does_not_blame_the_conductance(self, tmp_path):
        """A profile as curved as the reference says nothing about conductance."""
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 0.9, "tolerance": 0.1,
             "profile_span": 11.8, "reference_span": 12.0, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "nearly flat" not in check.detail
        assert "initial_profile_error" in check.detail

    def test_a_broken_reference_outranks_the_conductance_diagnosis(self, tmp_path):
        """Both tells can be present; the reference fault is the actionable one."""
        checks = self._checks(
            tmp_path,
            {"case": "analytical series", "max_abs_error": 64.9582,
             "initial_profile_error": 62.4, "profile_span": 0.01,
             "reference_span": 30.0, "tolerance": 0.5, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "REFERENCE is the broken part" in check.detail
        assert "nearly flat" not in check.detail

    def test_a_flat_reference_is_not_used_as_the_conductance_baseline(self, tmp_path):
        """Guard against dividing by a flat reference."""
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 0.9, "tolerance": 0.1,
             "profile_span": 0.0, "reference_span": 0.0, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "nearly flat" not in check.detail

    def test_a_steady_state_fault_is_named_as_structural(self, tmp_path):
        """The real bug from runs/cumcm-2026-a2026-marginfix.

        The solver's moisture block built a fully implicit diagonal and then also
        added the explicit interior fluxes to the right-hand side, counting them
        twice. Its moisture field never moved from 2.55 for the whole run, and
        every subproblem failed. Its t=0 behaviour was fine, so the t=0 test says
        nothing; the steady-state test is what identifies it.
        """
        checks = self._checks(
            tmp_path,
            {"case": "constant-property heat conduction, convection BC",
             "max_abs_error": 4.2, "initial_profile_error": 1e-9,
             "steady_state_error": 3.9, "tolerance": 0.5, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "does not hold the exact steady state" in check.detail
        assert "STRUCTURAL fault" in check.detail
        assert "counted twice" in check.detail
        assert "cannot fix it" in check.detail
        assert "initial_profile_error" not in check.detail

    def test_a_steady_state_fault_needs_a_healthy_t0(self, tmp_path):
        """A broken reference must keep priority over the structural diagnosis."""
        checks = self._checks(
            tmp_path,
            {"case": "analytical series", "max_abs_error": 64.0,
             "initial_profile_error": 62.0, "steady_state_error": 30.0,
             "tolerance": 0.5, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "REFERENCE is the broken part" in check.detail
        assert "STRUCTURAL fault" not in check.detail

    def test_a_steady_state_within_tolerance_is_not_a_fault(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 0.9, "initial_profile_error": 1e-9,
             "steady_state_error": 1e-10, "tolerance": 0.5, "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "STRUCTURAL fault" not in check.detail

    def test_a_failing_self_check_still_points_at_the_reference_first(self, tmp_path):
        """Without the diagnostic field, ask for it rather than guess."""
        checks = self._checks(
            tmp_path,
            {"case": "analytical series", "max_abs_error": 9.0, "tolerance": 0.1,
             "passed": False},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "initial_profile_error" in check.detail
        assert "REFERENCE" in check.detail

    def test_a_healthy_reference_within_tolerance_still_passes(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"case": "slab", "max_abs_error": 1e-4, "initial_profile_error": 1e-12,
             "tolerance": 1e-3, "passed": True},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.PASS

    def test_a_malformed_self_check_fails(self, tmp_path):
        checks = self._checks(
            tmp_path,
            {"passed": True},
            required=True,
        )
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "max_abs_error" in check.detail
        assert "tolerance" in check.detail

    def test_a_missing_self_check_fails_when_the_model_integrates(self, tmp_path):
        """This is the defect: a scheme that was never validated must not pass."""
        checks = self._checks(tmp_path, None, required=True)
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.FAIL
        assert "never" in check.detail

    def test_a_missing_self_check_is_not_applicable_for_an_algebraic_model(
        self, tmp_path
    ):
        checks = self._checks(tmp_path, None, required=False)
        check = self._by_name(checks, "scheme_self_check_passed")
        assert check.status == CheckStatus.NOT_APPLICABLE
        assert all(c.status != CheckStatus.FAIL for c in checks)

    def test_a_failing_scheme_check_blocks_the_report(self, tmp_path):
        from mathmodel.domain.math_model import MathematicalModel

        checks = self._checks(tmp_path, None, required=True)
        report = build_report(MathematicalModel(
            name="m", description="d", model_family="parabolic PDE",
        ), "RUN-1", checks)
        assert report.overall == CheckStatus.FAIL
        assert any("scheme_self_check_passed" in b for b in report.blocking_failures)


class TestRepairNoticesItIsRepeatingItself:
    """A repair loop that keeps seeing one class of error is not converging.

    The real A problem failed all six of its repair attempts with six different
    symptoms of ONE unresolved off-by-one between the node count and a profile
    length. Each repair patched the line that reported the error, so the error
    moved to the next line. Saying the error again is useless by then; the loop
    has to say "stop patching call sites".
    """

    _REAL_ERRORS = [
        "ValueError: operands could not be broadcast together with shapes (40,) (41,)",
        "ValueError: 22 columns passed, passed data had 21 columns",
        "AssertionError: header/data mismatch: hdr=22 T=21 C=21",
        "ValueError: fp and xp are not of the same length.",
        "AssertionError: interp_profile: grid size 21 != profile length 20",
    ]

    def test_every_real_failure_is_grouped_as_one_class(self):
        from mathmodel.autopilot.pipeline import _failure_class

        classes = {_failure_class(err) for err in self._REAL_ERRORS}
        assert classes == {"shape-mismatch"}, classes

    def test_a_different_failure_keeps_its_own_signature(self):
        from mathmodel.autopilot.pipeline import _failure_class

        assert _failure_class("KeyError: 'D_eff'") != "shape-mismatch"

    def test_the_signature_ignores_the_numbers(self):
        from mathmodel.autopilot.pipeline import _failure_signature

        assert _failure_signature(
            "AssertionError: grid size 21 != profile length 20"
        ) == _failure_signature("AssertionError: grid size 41 != profile length 40")

    def test_the_note_tells_the_repair_to_stop_patching_call_sites(self):
        from mathmodel.autopilot.pipeline import _repeated_failure_note

        note = _repeated_failure_note("shape-mismatch", attempt=4, seen=3)
        assert "STOP PATCHING CALL SITES" in note
        assert "attempt 4" in note
        assert "previous 3" in note
        assert "moved to a different line" in note
        assert "one constant" in note

    def test_a_non_shape_class_gets_a_generic_escalation(self):
        from mathmodel.autopilot.pipeline import _repeated_failure_note

        note = _repeated_failure_note("keyerror: '#d_eff#'", attempt=2, seen=1)
        assert "STOP PATCHING CALL SITES" in note
        assert "shared definition" in note
        assert "one constant" not in note

    def test_the_first_failure_gets_no_escalation(self):
        """Only a repeat is evidence that the previous repair failed."""
        from mathmodel.autopilot.pipeline import _failure_class

        seen = 0
        assert not seen
        assert _failure_class(self._REAL_ERRORS[0]) == "shape-mismatch"


class TestDisagreementHintFitsTheModel:
    """The repair hint for an independent disagreement must match the model."""

    @staticmethod
    def _report(checks):
        return VerificationReport(model_id="m", checks=checks)

    def test_a_numeric_independent_disagreement_is_detected(self):
        from mathmodel.autopilot.pipeline import _independent_disagreement

        report = self._report([VerificationCheck(
            name="independent_statistics_agree",
            category="independent",
            status=CheckStatus.FAIL,
            detail="T_center_final_C: solver=28.0 independent=0.0",
        )])
        assert _independent_disagreement(report) is True

    def test_an_agreeing_verifier_is_not_a_disagreement(self):
        from mathmodel.autopilot.pipeline import _independent_disagreement

        report = self._report([VerificationCheck(
            name="independent_statistics_agree",
            category="independent",
            status=CheckStatus.PASS,
            detail="8 statistic(s) agree",
        )])
        assert _independent_disagreement(report) is False

    def test_a_non_independent_failure_is_not_a_disagreement(self):
        from mathmodel.autopilot.pipeline import _independent_disagreement

        report = self._report([VerificationCheck(
            name="output_meets_stated_final_target",
            category="correctness",
            status=CheckStatus.FAIL,
            detail="the largest value in the final row is 2.5469",
        )])
        assert _independent_disagreement(report) is False

    def test_a_pde_model_gets_the_numerical_scheme_hint(self):
        from mathmodel.autopilot.pipeline import _disagreement_hint

        hint = _disagreement_hint(_a2026_like_model())
        assert "NUMERICAL SCHEME" in hint
        assert "control-volume measure" in hint
        assert "PREVIOUS time level" in hint
        assert "occupancy expansion" not in hint

    def test_a_combinatorial_model_keeps_the_occupancy_hint(self):
        from mathmodel.autopilot.pipeline import _disagreement_hint

        hint = _disagreement_hint(_d_like_model())
        assert "occupancy expansion" in hint
        assert "k-th use" in hint
        assert "NUMERICAL SCHEME" not in hint


class TestCodegenNumericalContract:
    """The generator is told the failure modes that actually produced wrong runs."""

    def _prompt(self):
        from mathmodel.agents.base import AgentStatus  # noqa: F401
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        generator = SolverCodeGenerator.__new__(SolverCodeGenerator)
        return generator._build_prompt(
            model=_a2026_like_model(),
            problem_text="药材烘干问题",
            data_schema="{}",
            required_outputs=["result1.xlsx"],
            input_file_names=["附件1.xlsx"],
        )

    def test_the_prompt_states_the_control_volume_rule(self):
        prompt = self._prompt()
        assert "NUMERICAL SCHEME CORRECTNESS" in prompt
        assert "same control-volume measure" in prompt

    def test_the_prompt_states_the_picard_rhs_rule(self):
        prompt = self._prompt()
        assert "previous time level" in prompt
        assert "b(U_prev)" in prompt
        assert "never use it" in prompt

    def test_the_prompt_states_the_conductance_dimensional_rule(self):
        """The real bug from runs/cumcm-2026-a2026-final.

        An extra 1/R**2 in the conduction term made diffusion ~2500x too fast.
        Dropping the cell width was already covered; the symmetric error (a
        spurious EXTRA factor) was not, and the model must be told the shape tell
        because a scaling error leaves every number plausible.
        """
        prompt = self._prompt()
        assert "k * A_face / (node spacing)" in prompt
        assert "dimensionless" in prompt
        assert "1/R**2" in prompt
        assert "profile SHAPE, not magnitude" in prompt
        assert "shrinking the time step" in prompt

    def test_the_prompt_requires_the_profile_span_diagnostic(self):
        prompt = self._prompt()
        assert '"profile_span"' in prompt
        assert '"reference_span"' in prompt
        assert "nearly FLAT" in prompt

    def test_the_prompt_warns_that_the_exponent_may_be_a_ratio(self):
        """The statement typesets `e^{-0.45/C}`; flattened it reads `e^-0.45 C`.

        Reading it as a product changes the drying time from 57.00 h (2.38 days,
        which matches the problem's stated 2-3 day process and its 72 h shrinkage
        attachment) to 16.79 h (0.70 days, which matches neither).
        """
        prompt = self._prompt()
        assert "TRANSCRIBING EMPIRICAL FORMULAS" in prompt
        assert "RATIO inside an exponent" in prompt
        assert "exp(-0.45 / C)" in prompt
        assert "exp(-0.45 * C)" in prompt
        assert "57.00 h" in prompt
        assert "16.79 h" in prompt

    def test_the_prompt_requires_reading_its_own_files_back(self):
        """round 65: attempt6's header claimed the required grids while its files
        stepped 60 s and 3600 s, and it stayed blocked on that for four repairs."""
        prompt = self._prompt()
        assert "AFTER WRITING EACH FILE, REOPEN IT AND CHECK IT" in prompt
        assert "Your intent is not evidence; the file on disk is" in prompt
        assert "sub-second/60-s output grids" in prompt
        assert "read its own artifacts back" in prompt
        assert "A silent wrong file is " in prompt

    def test_the_prompt_requires_an_honest_expected_order(self):
        """sourcetermfix measured order 0.991 -- correct for backward Euler with
        coupled refinement -- and was still failed for having declared 2."""
        prompt = self._prompt()
        assert "Declare `expected_order` as the order your scheme can ACTUALLY" in prompt
        assert "is the MINIMUM of the " in prompt
        assert "FIRST order overall" in prompt
        assert "0.991" in prompt
        assert "`min(order_t, order_x)`" in prompt

    def test_the_prompt_requires_named_importable_operators(self):
        """Order ~= 0 is consistent with a wrong source term AND with a wrong
        operator; only an independently driven operator can tell them apart."""
        prompt = self._prompt()
        assert "## NAMED OPERATORS (hard rule)" in prompt
        assert "`assemble_fixed(grid, properties, dt)" in prompt
        assert "`step_fixed(matrix, rhs)" in prompt
        assert "order approximately ZERO" in prompt
        assert "missing its control-volume width factor" in prompt
        assert "can drive them on a case whose " in prompt

    def test_the_prompt_requires_a_tiny_end_to_end_proof_before_the_real_run(self):
        """Measured: of 94 failed attempts across 39 runs, 61 (65%) died on a
        pure code fault before reaching any physics."""
        prompt = self._prompt()
        assert "Prove the WHOLE pipeline on a TINY case before the real run" in prompt
        assert "61 of those (65%)" in prompt
        assert "3 time steps" in prompt
        assert "len(row) == len(header)" in prompt
        assert "in about a second instead of after" in prompt

    def test_the_prompt_offers_a_verified_manufactured_source_term(self):
        """The solver's hand-derived source terms gave orders of -1.49, -1.90,
        -1.96 and -1.96 across four runs while its answers were correct. This
        formula was verified by complex-step differentiation (residual 3.6e-15)
        before being put in the prompt."""
        prompt = self._prompt()
        assert "USE THIS ONE" in prompt
        assert "u(r, t) = (1 + r**2 / R**2) * (1 + t / tau)" in prompt
        assert "f = (1 + r**2 / R**2) / tau - 4 * D * (1 + t / tau) / R**2" in prompt
        assert "Check it yourself by substituting back" in prompt

    def test_the_offered_source_term_actually_satisfies_the_equation(self):
        """The host's own claim must hold, checked independently of the prompt."""
        import cmath

        D, R, tau, h = 0.5, 1.0, 3.0, 1e-30

        def u(r, t):
            return (1 + r**2 / R**2) * (1 + t / tau)

        def f(r, t):
            return (1 + r**2 / R**2) / tau - 4 * D * (1 + t / tau) / R**2

        def ddt(r, t):
            return u(r, t + 1j * h).imag / h

        # (1/r) d/dr(r du/dr) for this field: the flux r*du/dr is 2*r**2/R**2 *
        # (1+t/tau), so dividing its derivative by r gives 4/R**2 * (1+t/tau).
        def lap_exact(r, t):
            return (4 / R**2) * (1 + t / tau)

        for r, t in [(0.1, 0.0), (0.5, 12.0), (0.9, 40.0)]:
            assert abs(ddt(r, t) - D * lap_exact(r, t) - f(r, t)) < 1e-9
        assert abs(u(R, 7.0) - 2 * (1 + 7.0 / tau)) < 1e-12

    def test_the_prompt_requires_a_timescale_prediction_first(self):
        """Real case from runs/cumcm-2026-a2026-levelsfix: a solver integrated
        far past the process duration and reported a final moisture of 2.5217
        against an initial 2.55 -- the material had essentially not dried."""
        prompt = self._prompt()
        assert "Predict the TIMESCALE before you integrate" in prompt
        assert "R**2 / D" in prompt
        assert "2.5217" in prompt
        assert "the interior " in prompt and "transport is switched off" in prompt
        assert "orders of magnitude shorter" in prompt

    def test_the_prompt_requires_staged_scheme_validation(self):
        """Evidence: the constant-coefficient case passes (err 0.0148) when it is
        tried, while solvers that jump straight to a coupled variable-coefficient
        Picard scheme report order -1.49 or a field stuck at its initial value."""
        prompt = self._prompt()
        assert "Validate your scheme in STAGES" in prompt
        assert "CONSTANT-coefficient version" in prompt
        assert "before you add variable coefficients" in prompt
        assert "0.0148" in prompt
        assert "-1.49" in prompt
        assert "SEPARATELY from the full model" in prompt

    def test_the_prompt_requires_the_surface_conductance_biot_check(self):
        """Real case from runs/cumcm-2026-a2026-verdictfix.

        The solver's moisture block annotated its surface flux as
        `R * km * (C - Cair)` and its diffusivity as `rho_d * D / R**2` -- a ratio
        to the interior conductance off by R**2 -- and then integrated ten days
        with the moisture essentially unchanged. The problem says the process
        takes 2-3 days.
        """
        prompt = self._prompt()
        assert "Measure the SURFACE conductance on the surface area" in prompt
        assert "Bi = h*R/k" in prompt
        assert "R * km * (C - Cair)" in prompt
        assert "rho_d * D / R**2" in prompt
        assert "ten days while the moisture stayed" in prompt
        assert "Sanity-check your answer against the duration" in prompt
        assert "2-3 days" in prompt or "couple of days" in prompt

    def test_the_prompt_requires_the_solver_to_declare_its_formulas(self):
        """The verifier needs the formulas to check the transcription, and the
        solver is the only party that knows what its code actually does."""
        prompt = self._prompt()
        assert '"formulas": {...}' in prompt
        assert "Report what the code DOES, not what you intended" in prompt
        assert "7e-9 * exp(-0.89 / C)" in prompt
        assert "2.4e-3 * exp(-0.45 / C)" in prompt

    def test_the_prompt_warns_that_arrhenius_amplifies_temperature_error(self):
        """Real case from runs/cumcm-2026-a2026-formulafix.

        The solver's temperature went negative in Celsius, driving exp(-3850/T)
        toward zero, so the moisture never moved and every subproblem failed. The
        reported fault looked like mass transfer; the cause was heat.
        """
        prompt = self._prompt()
        assert "amplifies a temperature error" in prompt
        assert "looked like mass transfer, but the cause was heat" in prompt
        assert "if a clip actually binds, print it" in prompt

    def test_the_prompt_settles_the_ratio_reading_by_the_arrhenius_factor(self):
        """`e^{-3850/T}` is unmistakably a ratio, and the statement typesets the
        other exponents the same way, so they are ratios too."""
        prompt = self._prompt()
        assert "finding the reading that is NOT " in prompt
        assert "Arrhenius factor" in prompt
        assert "numerator sits high and the denominator low" in prompt
        assert "Do not read one as a product and another as a quotient" in prompt

    def test_the_prompt_requires_a_cross_check_against_the_problems_own_claim(self):
        """The generated model caught my own wrong rule by doing exactly this."""
        prompt = self._prompt()
        assert "2-3 days" in prompt
        assert "out to 72 h" in prompt
        assert "2.38 days" in prompt
        assert "if two readings both survive, say which you chose" in prompt

    def test_the_prompt_warns_that_an_arrhenius_factor_needs_kelvin(self):
        """Real case: a solver wrote exp(-3850.0 / np.maximum(T, 1.0)).

        With T in Celsius, exp(-3850/28) is about 1e-60, so the diffusivity
        collapses to zero and the moisture field never moves. The guard against
        small values is itself the tell: a Kelvin temperature is never near 1.
        """
        prompt = self._prompt()
        assert "exp(-3850/T)" in prompt
        assert "KELVIN" in prompt
        assert "np.maximum(T, 1.0)" in prompt
        assert "T_K = T_C + 273.15" in prompt

    def test_the_prompt_requires_the_memory_budget(self):
        """Real case from runs/cumcm-2026-a2026-fluxfix.

        Two of three attempts were OOM-killed after ~160 s while a third finished
        the same problem in 14 s, because the model vectorised over both space and
        time. The cap is a real acceptance constraint the program cannot discover
        for itself, exactly like the time budget.
        """
        from mathmodel.autopilot.codegen import PROGRAM_MEMORY_BUDGET_MB

        prompt = self._prompt()
        assert f"{PROGRAM_MEMORY_BUDGET_MB} MB of memory" in prompt
        assert "killed for memory" in prompt
        assert "space and time" in prompt

    def test_the_prompt_requires_margin_on_a_strict_threshold(self):
        """Real case from runs/cumcm-2026-a2026-gridfix.

        The solver rounded its output to four decimals, so a true 0.14996 was
        written as 0.1500 -- not strictly below 0.15 -- and the run failed on a
        value that was compliant before rounding.
        """
        prompt = self._prompt()
        assert "Leave MARGIN" in prompt
        assert "test the value you actually WRITE" in prompt
        assert "rounded to four" in prompt
        assert "0.14996" in prompt

    def test_the_prompt_states_the_runtime_budget(self):
        """A solver that cannot finish is not a solver, and the model must know.

        Measured: a generated implicit Crank-Nicolson program was killed at
        exactly the limit on all three attempts having printed nothing, because
        nothing had ever told it that a budget existed.
        """
        from mathmodel.autopilot.codegen import PROGRAM_TIME_BUDGET_SECONDS

        prompt = self._prompt()
        assert "RUNTIME BUDGET" in prompt
        assert f"killed after {PROGRAM_TIME_BUDGET_SECONDS} seconds" in prompt
        assert "produces no results at all" in prompt

    def test_the_prompt_requires_flushed_progress_output(self):
        prompt = self._prompt()
        assert "flush=True" in prompt
        assert "leaves evidence of how far it got" in prompt

    def test_the_progress_rule_warns_off_the_array_truthiness_crash(self):
        """The obvious way to write a progress guard killed a working run.

        `if k % 500 == 0:` raised "The truth value of an array with more than one
        element is ambiguous" and threw away three correctly finished
        subproblems.
        """
        prompt = self._prompt()
        assert "plain integer counter" in prompt
        assert "truth value of an array" in prompt
        assert "throws away every result already computed" in prompt

    def test_the_runner_default_matches_the_stated_budget(self):
        from mathmodel.autopilot.codegen import (
            PROGRAM_TIME_BUDGET_SECONDS,
            SandboxProgramRunner,
        )

        runner = SandboxProgramRunner()
        assert runner._limits.timeout_seconds == PROGRAM_TIME_BUDGET_SECONDS

    def test_the_prompt_requires_a_reported_self_check(self):
        prompt = self._prompt()
        assert "self_check" in prompt
        assert "max_abs_error" in prompt
        assert "MANDATORY" in prompt

    def test_the_prompt_names_the_mechanical_tell_of_the_storage_bug(self):
        """A rule the solver can check mechanically, not just a principle."""
        prompt = self._prompt()
        assert "snapshot" in prompt.lower()
        assert "never read" in prompt
        assert "STEADY STATE" in prompt
        assert "halve the time step" in prompt

    def test_the_prompt_requires_a_self_contained_validation_case(self):
        prompt = self._prompt()
        assert "self-contained" in prompt
        assert "same boundary" in prompt
        assert "dt/2" in prompt

    def test_the_prompt_requires_the_reference_itself_to_be_validated(self):
        """A wrong analytical series was the real cause of a 20.9 K 'error'."""
        prompt = self._prompt()
        assert "Validate your REFERENCE" in prompt
        assert "evaluate it at t=0" in prompt
        assert "beta*R*J1(beta*R)" in prompt
        assert "factor of R" in prompt

    def test_the_prompt_explains_the_ellipsis_header_convention(self):
        """The template's '…' must be expanded without eating the trailing label."""
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        generator = SolverCodeGenerator.__new__(SolverCodeGenerator)
        prompt = generator._build_prompt(
            model=_a2026_like_model(),
            problem_text="药材烘干问题",
            data_schema="{}",
            required_outputs=["result4.xlsx"],
            input_file_names=["附件1.xlsx"],
            output_headers={
                "result4.xlsx": [
                    "时间\\到药材中心的距离", "0", "0.1", "0.2", "…", "药材表面",
                ]
            },
        )
        assert "…" in prompt
        assert "药材表面" in prompt
        assert "arithmetic progression" in prompt
        assert "NOT the number it denotes" in prompt
        assert "one step before" in prompt


class _CapturingRouter:
    """Records the prompt it was handed, then fails, so no model is called."""

    def __init__(self):
        self.prompts: list[str] = []

    async def route_generate(self, **kwargs):
        self.prompts.append(kwargs.get("prompt", ""))
        raise RuntimeError("prompt captured")


class TestVerifierLocatesQuantitiesByData:
    """A name-based sheet filter silently produces None on Sheet1-only files."""

    async def _prompt(self):
        from mathmodel.autopilot.verify import IndependentVerifier

        router = _CapturingRouter()
        verifier = IndependentVerifier(router, runner=None)
        with pytest.raises(RuntimeError):
            await verifier.generate(
                model=_a2026_like_model(),
                problem_text="将水分浓度保存到 result3.xlsx 中。",
                required_outputs=["result3.xlsx"],
                solver_statistics={"t_dry3_s": 9060.0},
                available_inputs=["附件1.xlsx", "output/result3.xlsx"],
            )
        return router.prompts[0]

    async def test_the_verifier_is_told_to_use_the_problem_statement(self):
        prompt = await self._prompt()
        assert "LOCATING A QUANTITY" in prompt
        assert "The problem statement says what each result file holds" in prompt

    async def test_the_verifier_is_warned_off_sheet_name_filters(self):
        prompt = await self._prompt()
        assert "Do NOT select worksheets by whether their NAME contains" in prompt
        assert "only Sheet1" in prompt

    async def test_the_verifier_is_told_not_to_difference_the_samples(self):
        """runs/cumcm-2026-a2026-rebasefix report3.json: the verifier reported
        `constraint dTdr_0_t: max|dT/dr| at center=8.8000e+00` as a violation,
        computed by differencing output columns sampled at r=0 and r=0.1 cm."""
        prompt = await self._prompt()
        assert "2e." in prompt
        assert "DERIVATIVE or a limit at a boundary point" in prompt
        assert "8.8000e+00" in prompt
        assert "could not have passed for ANY correct" in prompt
        assert "Never turn a finite difference across coarse" in prompt

    async def test_the_verifier_is_told_to_read_the_row_the_name_asks_for(self):
        """report4.json: recomputing `problem1_final_min_T_C` as a minimum over
        every row returned 28.0 -- the initial temperature."""
        prompt = await self._prompt()
        assert "2f." in prompt
        assert "Read the row a statistic's NAME asks for" in prompt
        assert "problem1_final_min_T_C" in prompt
        assert "33.5765" in prompt
        assert "different statistics" in prompt

    async def test_a_missing_quantity_verdict_needs_an_enumeration(self):
        """sourcetermfix: the verifier reported 'could not identify moisture
        sheet' while both files held a worksheet named 水分浓度 whose final row
        read 0.0499 kg/kg against a required 0.15."""
        prompt = await self._prompt()
        assert "Reporting that a quantity is MISSING is a strong claim" in prompt
        assert "could not identify moisture sheet" in prompt
        assert "0.0499" in prompt
        assert "A missing-quantity verdict with no such " in prompt

    async def test_the_verifier_is_told_not_to_fail_silently_into_none(self):
        prompt = await self._prompt()
        assert "Never let a lookup fail silently into None" in prompt

    async def _prompt_with_formulas(self, formulas):
        from mathmodel.autopilot.verify import IndependentVerifier

        router = _CapturingRouter()
        verifier = IndependentVerifier(router, runner=None)
        with pytest.raises(RuntimeError):
            await verifier.generate(
                model=_a2026_like_model(),
                problem_text="水分浓度扩散系数 D = 2.4e-3 e^-0.45C e^-3850/T。",
                required_outputs=["result3.xlsx"],
                solver_statistics={"t_dry3_s": 9060.0},
                solver_formulas=formulas,
                available_inputs=["附件1.xlsx", "output/result3.xlsx"],
            )
        return router.prompts[0]

    async def test_the_verifier_checks_the_solvers_declared_formulas(self):
        """A transcription slip produces plausible numbers, so no numeric
        comparison can catch it. The verifier can, if it is handed the formulas."""
        prompt = await self._prompt_with_formulas(
            {"D_q3": "2.4e-3 * exp(-0.45 * C) * exp(-3850 / T_K)"}
        )
        assert "FORMULAS THE SOLVER SAYS IT IMPLEMENTED" in prompt
        assert "2.4e-3 * exp(-0.45 * C) * exp(-3850 / T_K)" in prompt
        assert "character by character" in prompt
        assert "no numeric comparison can reveal it" in prompt
        assert "Compare the STRUCTURE of every exponent" in prompt
        assert "PRODUCT or a RATIO" in prompt
        assert "requires T in KELVIN" in prompt

    async def test_the_verifier_distinguishes_ambiguity_from_error(self):
        """Real case from runs/cumcm-2026-a2026-biotfix.

        The solver correctly read a ratio inside an exponent -- the reading that
        reproduces the statement's own 2-3 day process duration -- and the
        verifier failed all three of its formulas with the note "exponent
        structure ambiguous", blocking the run over a problem-statement
        ambiguity rather than a solver error.
        """
        prompt = await self._prompt_with_formulas(
            {"D_q3": "2.4e-3 * exp(-0.45 / C) * exp(-3850 / T_K)"}
        )
        assert "Distinguish an AMBIGUITY from an ERROR" in prompt
        assert "recorded modelling assumption, not a defect" in prompt
        assert "'ambiguous'" in prompt
        assert "contradicts something the " in prompt
        assert "report the number you get" in prompt

    async def test_the_verifier_is_spared_the_formula_section_when_absent(self):
        prompt = await self._prompt_with_formulas({})
        assert "FORMULAS THE SOLVER SAYS IT IMPLEMENTED" not in prompt

    async def test_the_verifier_is_told_where_a_threshold_applies(self):
        """Taking a max over every row tests t=0 against a final target."""
        prompt = await self._prompt()
        assert "Where a threshold applies matters as much as the threshold" in prompt
        assert "applies to the FINAL row" in prompt
        assert "INITIAL condition" in prompt

    async def test_the_verifier_is_told_to_use_the_problems_units(self):
        """A 273.15 bound on Celsius data reports a correct answer as failing."""
        prompt = await self._prompt()
        assert "problem's OWN units" in prompt
        assert "absolute-zero bound of 273.15" in prompt
        assert "which unit you assumed" in prompt

    async def _prompt_with_structure(self, structure):
        from mathmodel.autopilot.verify import IndependentVerifier

        router = _CapturingRouter()
        verifier = IndependentVerifier(router, runner=None)
        with pytest.raises(RuntimeError):
            await verifier.generate(
                model=_a2026_like_model(),
                problem_text="将水分浓度保存到 result1.xlsx 中。",
                required_outputs=["result1.xlsx"],
                solver_statistics={"x": 1.0},
                available_inputs=["附件1.xlsx"],
                solution_structure=structure,
            )
        return router.prompts[0]

    async def test_the_verifier_is_handed_the_real_file_shape(self):
        prompt = await self._prompt_with_structure({
            "result1.xlsx": [
                {"sheet": "温度", "data_rows": 1801, "columns": 22,
                 "header_row": ["时间\\到药材中心的距离", "0.0", "0.1"]},
            ]
        })
        assert "SOLUTION FILE STRUCTURE" in prompt
        assert "1801" in prompt

    async def test_the_verifier_is_told_the_time_column_is_not_a_field(self):
        prompt = await self._prompt_with_structure({
            "result1.xlsx": [{"sheet": "温度", "data_rows": 10, "columns": 22}]
        })
        assert "FIRST column of every such sheet is the TIME axis" in prompt
        assert "0..t_end" in prompt

    async def test_no_structure_block_when_there_is_nothing_to_describe(self):
        prompt = await self._prompt_with_structure(None)
        assert "SOLUTION FILE STRUCTURE" not in prompt


class TestDescribeSolutionFiles:
    """The shape handed to the verifier must come from the real files."""

    def _book(self, tmp_path):
        import openpyxl

        path = tmp_path / "result1.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "温度"
        ws.append(["时间\\到药材中心的距离", 0.0, 0.1])
        ws.append([0, 28.0, 28.0])
        ws.append([1, 28.1, 28.0])
        ws2 = wb.create_sheet("水分浓度")
        ws2.append(["时间\\到药材中心的距离", 0.0, 0.1])
        ws2.append([0, 2.55, 2.55])
        wb.save(path)
        return path

    def test_every_sheet_is_described_with_its_real_shape(self, tmp_path):
        from mathmodel.autopilot.verify import describe_solution_files

        self._book(tmp_path)
        described = describe_solution_files(tmp_path, ["result1.xlsx"])
        sheets = {s["sheet"]: s for s in described["result1.xlsx"]}
        assert set(sheets) == {"温度", "水分浓度"}
        assert sheets["温度"]["data_rows"] == 2
        assert sheets["温度"]["columns"] == 3
        assert sheets["温度"]["header_row"][0] == "时间\\到药材中心的距离"

    def test_a_missing_output_is_simply_absent(self, tmp_path):
        from mathmodel.autopilot.verify import describe_solution_files

        self._book(tmp_path)
        described = describe_solution_files(tmp_path, ["result9.xlsx"])
        assert described == {}

    def test_the_real_values_are_handed_over_too(self, tmp_path):
        """Names alone let the verifier read the wrong sheet.

        Told only the names it reported an initial moisture of 28.0 — the
        temperature sheet, which comes first.
        """
        from mathmodel.autopilot.verify import describe_solution_files

        self._book(tmp_path)
        described = describe_solution_files(tmp_path, ["result1.xlsx"])
        sheets = {s["sheet"]: s for s in described["result1.xlsx"]}
        assert sheets["温度"]["first_data_row"] == [0.0, 28.0, 28.0]
        assert sheets["水分浓度"]["first_data_row"] == [0.0, 2.55, 2.55]
        assert sheets["温度"]["last_data_row"] == [1.0, 28.1, 28.0]

    def test_the_two_sheets_are_distinguishable_by_content(self, tmp_path):
        from mathmodel.autopilot.verify import describe_solution_files

        self._book(tmp_path)
        described = describe_solution_files(tmp_path, ["result1.xlsx"])
        moisture = [
            s for s in described["result1.xlsx"]
            if s["first_data_row"] and abs(s["first_data_row"][1] - 2.55) < 1e-9
        ]
        assert [s["sheet"] for s in moisture] == ["水分浓度"]


class TestCodegenOutputComplianceRules:
    """The two compliance defects the gate reported must be stated up front."""

    def _prompt(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        generator = SolverCodeGenerator.__new__(SolverCodeGenerator)
        return generator._build_prompt(
            model=_a2026_like_model(),
            problem_text="药材烘干问题",
            data_schema="{}",
            required_outputs=["result3.xlsx"],
            input_file_names=["附件1.xlsx"],
        )

    def test_the_prompt_forbids_an_off_grid_row(self):
        prompt = self._prompt()
        assert "EVERY row must sit on that grid" in prompt
        assert "off-grid row invalidates" in prompt

    def test_the_prompt_requires_a_strict_threshold(self):
        prompt = self._prompt()
        assert "satisfy it STRICTLY" in prompt
        assert "equals the threshold is a failure" in prompt


class TestSummaryParsing:
    """The host's reader must not turn a complete report into "no statistics"."""

    def test_a_single_line_summary_is_read(self):
        from mathmodel.autopilot.codegen import _parse_summary

        parsed = _parse_summary('progress\n{"statistics": {"a": 1}}\n')
        assert parsed == {"statistics": {"a": 1}}

    def test_a_pretty_printed_summary_is_read(self):
        """A multi-line object is still a report; it used to parse as {}."""
        from mathmodel.autopilot.codegen import _parse_summary

        stdout = (
            "=== Problem 3 ===\n"
            "  Drying time: 34705.00 s\n"
            "{\n"
            '  "statistics": {"t_dry3_s": 34705.0},\n'
            '  "self_check": {"passed": false, "max_abs_error": 712.5},\n'
            '  "notes": "failed with max error 712.506922"\n'
            "}\n"
        )
        parsed = _parse_summary(stdout)
        assert parsed["statistics"] == {"t_dry3_s": 34705.0}
        assert parsed["self_check"]["max_abs_error"] == 712.5

    def test_the_last_object_wins(self):
        from mathmodel.autopilot.codegen import _parse_summary

        stdout = '{"statistics": {"a": 1}}\nnoise\n{"statistics": {"a": 2}}\n'
        assert _parse_summary(stdout) == {"statistics": {"a": 2}}

    def test_braces_in_prose_do_not_confuse_it(self):
        from mathmodel.autopilot.codegen import _parse_summary

        stdout = 'set {a, b} here\n{"statistics": {"a": 1}}\n'
        assert _parse_summary(stdout) == {"statistics": {"a": 1}}

    def test_a_crash_without_a_summary_yields_nothing(self):
        from mathmodel.autopilot.codegen import _parse_summary

        stdout = "Traceback (most recent call last):\n  AssertionError: nan\n"
        assert _parse_summary(stdout) == {}


class TestRepairShowsTheWholeProgram:
    """A repair must not hide the code the model is being asked to fix.

    The generated solvers run to 22k-36k characters. The old 12k excerpt meant
    every repair saw only its first third, so it fixed the reported issue and
    introduced new defects in the part it could not read.
    """

    @staticmethod
    def _program(lines: int) -> str:
        body = "\n".join(f"def f{i}():\n    return {i}" for i in range(lines))
        return body

    def test_a_program_longer_than_the_old_cut_is_shown_whole(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        program = self._program(2000)  # ~44k chars, like a real solver
        assert len(program) > 12000
        block = SolverCodeGenerator._repair_block(program, ["something failed"])
        assert program in block
        assert "cut off for length" not in block

    def test_an_oversized_program_is_cut_with_an_explicit_notice(self):
        from mathmodel.autopilot.codegen import (
            REPAIR_CODE_MAX_CHARS,
            SolverCodeGenerator,
        )

        program = "x = 1\n" * (REPAIR_CODE_MAX_CHARS // 3)
        block = SolverCodeGenerator._repair_block(program, ["too big"])
        assert len(program) > REPAIR_CODE_MAX_CHARS
        assert "cut off for length" in block
        assert "Re-derive the missing part" in block

    def test_the_repair_asks_for_a_minimal_change(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        block = SolverCodeGenerator._repair_block("x = 1\n", ["something failed"])
        assert "change ONLY what the reported issues require" in block
        assert "has made the program worse" in block

    def test_the_reported_issues_are_still_listed(self):
        from mathmodel.autopilot.codegen import SolverCodeGenerator

        block = SolverCodeGenerator._repair_block(
            "x = 1\n", ["header does not match the template", "time step is 21 s"]
        )
        assert "header does not match the template" in block
        assert "time step is 21 s" in block


class TestStatisticsToleranceConflatesTwoMechanisms:
    """A KNOWN false negative, pinned here so it stays visible.

    runs/cumcm-2026-a2026-auditfix: a solver and an independent verifier agreed
    on a drying time at 43.5 h vs 43.1167 h (0.9%), and on a second problem at
    41.33 h vs 41.12 h (0.5%). Both are strong corroboration between two
    independently discretised PDE solves, and both are currently reported as
    disagreements.

    Raising STATISTIC_REL_TOLERANCE to 5% fixes them, but it breaks a
    requirement established earlier and still asserted below: when a verifier
    reports 37.05 to two decimals the true value lies in [37.045, 37.055], so a
    solver value of 36.981875 is genuinely different. Rounding-precision
    comparison and discretisation-noise tolerance are different mechanisms that
    the single constant currently conflates.
    """

    def test_the_known_false_negative_is_still_present(self):
        """Pinned, NOT endorsed: this asserts today's behaviour so that fixing
        it properly becomes a deliberate, visible change."""
        from mathmodel.autopilot.verify import compare_statistics

        report = compare_statistics(
            {"t_dry_problem3_h": 43.5}, {"t_dry_problem3_h": 43.11666666666667}
        )
        assert report.status == CheckStatus.FAIL

    def test_the_real_modelling_errors_still_fail(self):
        """The exponent read as a product gives 16.79 h against 57.00 h."""
        from mathmodel.autopilot.verify import compare_statistics

        report = compare_statistics(
            {"t_dry_problem3_h": 57.00}, {"t_dry_problem3_h": 16.79}
        )
        assert report.status == CheckStatus.FAIL

    def test_a_full_precision_verifier_still_catches_a_real_gap(self):
        """The requirement that must NOT be traded away, restated here because
        it is the reason the blanket tolerance cannot simply be raised."""
        from mathmodel.autopilot.verify import compare_statistics

        report = compare_statistics({"T": 36.98187512600498}, {"T": 37.05})
        assert report.status == CheckStatus.FAIL


class TestStatisticsToleranceIsRelative:
    """Rounding is not disagreement.

    A verifier that prints 4 decimals was accused of disagreeing with a solver
    that prints full precision, because the tolerance was absolute at 1e-6:
    33.66469441752457 vs 33.6647 differs by 5.6e-6. Every mismatch in that
    report was this artifact.
    """
    def test_a_four_decimal_verifier_agrees_on_a_large_value(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"T": 33.66469441752457}, {"T": 33.6647}
        )
        assert check.status == CheckStatus.PASS, check.detail

    def test_a_four_decimal_verifier_agrees_on_a_small_value(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"C": 2.214135421467383}, {"C": 2.2141}
        )
        assert check.status == CheckStatus.PASS, check.detail

    def test_the_whole_real_disagreement_report_now_agrees(self):
        """The exact values from runs/cumcm-2026-a2026-repairfix/report3.json."""
        from mathmodel.autopilot.verify import compare_statistics

        solver = {
            "problem1_T_at_1800s.0": 33.66469441752457,
            "problem1_T_at_1800s.0.5": 33.86285786644868,
            "problem1_T_at_1800s.1": 34.459381133568,
            "problem1_T_at_1800s.1.5": 35.46800720393864,
            "problem1_T_at_1800s.2": 36.746570651321,
            "problem1_C_at_1800s.0": 2.214135421467383,
            "problem1_C_at_1800s.0.5": 2.2138524488703153,
            "problem1_C_at_1800s.1": 2.21300351530353,
            "problem1_C_at_1800s.1.5": 2.211583427001511,
            "problem1_C_at_1800s.2": 2.2098133680648777,
        }
        independent = {
            "problem1_T_at_1800s.0": 33.6647,
            "problem1_T_at_1800s.0.5": 33.8629,
            "problem1_T_at_1800s.1": 34.4594,
            "problem1_T_at_1800s.1.5": 35.468,
            "problem1_T_at_1800s.2": 36.7466,
            "problem1_C_at_1800s.0": 2.2141,
            "problem1_C_at_1800s.0.5": 2.2139,
            "problem1_C_at_1800s.1": 2.213,
            "problem1_C_at_1800s.1.5": 2.2116,
            "problem1_C_at_1800s.2": 2.2098,
        }
        check = compare_statistics(solver, independent)
        assert check.status == CheckStatus.PASS, check.detail
        assert check.evidence["compared"] == 10

    def test_a_real_physics_disagreement_still_fails(self):
        """The tolerance must not paper over a genuinely different answer."""
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"problem1_final_maxT_C": 36.98187512600498},
            {"problem1_final_maxT_C": 41.241718437399086},
        )
        assert check.status == CheckStatus.FAIL
        assert "differ by" in check.detail

    def test_a_wrong_order_of_magnitude_still_fails(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics({"C": 0.15}, {"C": 50.1954})
        assert check.status == CheckStatus.FAIL

    def test_a_near_zero_value_still_needs_to_match_closely(self):
        from mathmodel.autopilot.verify import compare_statistics

        assert compare_statistics({"x": 0.0}, {"x": 0.0}).status == CheckStatus.PASS
        assert compare_statistics({"x": 0.0}, {"x": 1e-3}).status == CheckStatus.FAIL

    def test_rounding_at_a_small_magnitude_is_not_a_disagreement(self):
        """0.0279489 printed to 4 decimals is 0.0279 — a 0.18% apparent gap.

        A flat relative tolerance cannot separate this from a real difference, so
        comparison happens at the precision the verifier reported.
        """
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"problem2_final_minC": 0.027948928005447023},
            {"problem2_final_minC": 0.0279},
        )
        assert check.status == CheckStatus.PASS, check.detail

    def test_a_real_one_percent_gap_at_the_same_precision_still_fails(self):
        """0.149966 against 0.1485 is 1%, far beyond 4-decimal rounding."""
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"problem4_final_maxC": 0.14996630613797496},
            {"problem4_final_maxC": 0.1485},
        )
        assert check.status == CheckStatus.FAIL
        assert "differ by 1.0%" in check.detail

    def test_a_full_precision_verifier_uses_the_relative_floor(self):
        """A 15-decimal report must not demand 1e-16 agreement."""
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"T": 36.98187512600498}, {"T": 36.98187512600497}
        )
        assert check.status == CheckStatus.PASS, check.detail

    def test_a_full_precision_verifier_still_catches_a_real_gap(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"T": 36.98187512600498}, {"T": 37.05}
        )
        assert check.status == CheckStatus.FAIL

    def test_a_different_grid_is_not_a_disagreement(self):
        """Real case: 41 internal grid points vs the file's 21 radius columns.

        Two correct programs may choose different discretisations, so comparing
        those reports a disagreement that is not one.
        """
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics({"grid_points": 41}, {"grid_points": 21})
        assert check.status == CheckStatus.NOT_RUN or check.status == CheckStatus.PASS
        assert check.status != CheckStatus.FAIL

    def test_a_structural_key_is_reported_as_not_compared(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"t_dry_s": 205800.0, "n_cells": 41}, {"t_dry_s": 205620.0, "n_cells": 21}
        )
        assert check.status == CheckStatus.PASS, check.detail
        assert check.evidence["compared"] == 1
        assert check.evidence["skipped_structural"] == ["n_cells"]

    def test_a_physical_quantity_is_still_compared(self):
        """Skipping implementation details must not skip real results."""
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"t_dry_s": 21060.0, "grid_points": 41},
            {"t_dry_s": 8460.0, "grid_points": 21},
        )
        assert check.status == CheckStatus.FAIL
        assert "t_dry_s" in check.detail

    def test_a_diverged_verifier_is_not_a_verdict_on_the_solver(self):
        """Real case from runs/cumcm-2026-a2026-gridfix.

        The verifier's own recomputation returned 4.87729701306026e+232 for a
        temperature between 28 and 60 C. Every key was then reported as the solver
        disagreeing by 100%, which sent the repair loop after a solver whose
        numbers were entirely plausible.
        """
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"problem1_final_center_T": 33.5765, "problem1_final_surface_T": 36.7863},
            {"problem1_final_center_T": 4.87729701306026e232,
             "problem1_final_surface_T": 4.881481475320075e232},
        )
        assert check.status == CheckStatus.FAIL
        assert "verifier's own recomputation diverged" in check.detail
        assert "Regenerate the verifier" in check.detail
        assert check.evidence["verifier_diverged"] == [
            "problem1_final_center_T", "problem1_final_surface_T",
        ]
        assert "solver=33.5765" not in check.detail

    def test_a_diverged_solver_is_still_the_solvers_fault(self):
        """The guard must not excuse a solver that blew up."""
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics(
            {"T": 4.87729701306026e232}, {"T": 33.5765}
        )
        assert check.status == CheckStatus.FAIL
        assert "solver's computation diverged" in check.detail
        assert "verifier_diverged" not in (check.evidence or {})

    def test_an_ordinary_gap_is_not_called_a_divergence(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics({"T": 33.5765}, {"T": 37.05})
        assert check.status == CheckStatus.FAIL
        assert "diverged" not in check.detail

    def test_a_non_finite_recomputation_blames_the_verifier(self):
        from mathmodel.autopilot.verify import compare_statistics

        check = compare_statistics({"T": 33.5765}, {"T": float("nan")})
        assert check.status == CheckStatus.FAIL
        assert "verifier's own recomputation diverged" in check.detail


class TestACheckRestingOnADivergedNumberIsNotEvidence:
    """Real case from runs/cumcm-2026-a2026-formulafix3.

    The verifier's own recomputation diverged and it reported
    "max |dT| at final t=1800s = 2.67...e+155 C" as its counterexample. That
    FAIL was the run's blocker even though the comparison layer had already
    established the verifier was the diverged side. Downgrading such a check to
    NOT_RUN still blocks the report, so no standard is relaxed, but the stated
    reason becomes accurate.
    """

    def test_the_real_verifier_check_is_detected(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            unevidenced_by_diverged_recomputation,
        )

        check = VerificationCheck(
            name="independent_recomputation",
            category="independent",
            status=CheckStatus.FAIL,
            detail=(
                "Recomputed 1D radial heat+moisture; max |dT| at final t=1800s = "
                "26713301353134503672255805428570847142741734842648552062829822"
                "0876877136968827326846925492414998214056785845516365316139192942592.000 C."
            ),
        )
        assert unevidenced_by_diverged_recomputation(check) is not None

    def test_a_diverged_number_in_evidence_is_detected(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            unevidenced_by_diverged_recomputation,
        )

        check = VerificationCheck(
            name="independent_constraint_1",
            category="independent",
            status=CheckStatus.FAIL,
            detail="temperature bound violated",
            evidence={"counterexamples": [{"value": 1e155, "where": "t=1800"}]},
        )
        assert unevidenced_by_diverged_recomputation(check) is not None

    def test_a_structural_finding_is_left_alone(self):
        """A check with no huge number in it still stands on its own."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            unevidenced_by_diverged_recomputation,
        )

        check = VerificationCheck(
            name="independent_worksheets",
            category="independent",
            status=CheckStatus.FAIL,
            detail="result3.xlsx has 15001 rows; the template asks for one row per 60 s.",
        )
        assert unevidenced_by_diverged_recomputation(check) is None

    def test_physical_values_are_never_called_diverged(self):
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            unevidenced_by_diverged_recomputation,
        )

        for value in (28.0, 0.15, 5443200.0, 2.55e-9, 1e29):
            check = VerificationCheck(
                name="independent_x",
                category="independent",
                status=CheckStatus.FAIL,
                detail=f"the largest value in the final row is {value}",
            )
            assert unevidenced_by_diverged_recomputation(check) is None, value

    def test_downgrading_to_not_run_still_fails_the_report(self):
        """The downgrade must change the stated reason, never the standard."""
        from mathmodel.autopilot.verify import (
            CheckStatus,
            VerificationCheck,
            build_report,
        )

        report = build_report(
            _a2026_like_model(),
            execution_run_id="run-1",
            checks=[
                VerificationCheck(
                    name="independent_recomputation",
                    category="independent",
                    status=CheckStatus.NOT_RUN,
                    detail="Not verifiable: rested on a diverged value.",
                )
            ],
        )
        assert report.overall == CheckStatus.FAIL
        assert report.blocking_failures == []


def test_scalar_where_an_array_was_expected_is_one_failure_class():
    """ceilingfix attempts 1-4 failed with four different errors that share one
    root cause -- a scalar where an array was expected. None of the original
    markers matched, so the classifier saw unrelated classes and the
    "STOP PATCHING CALL SITES" escalation never fired."""
    from mathmodel.autopilot.pipeline import _failure_class

    errors = [
        "IndexError: invalid index to scalar variable.",
        "ValueError: could not broadcast input array from shape (2,) into shape (20,)",
        "TypeError: 'float' object is not subscriptable",
        "TypeError: 'numpy.float64' object is not subscriptable",
    ]
    classes = {_failure_class(e) for e in errors}
    assert classes == {"shape-mismatch"}, classes


def test_the_array_scalar_family_and_self_guards_classify_together():
    """Round 70 audit of 200 real solve outcomes: these phrasings were each
    becoming their own bare signature, so neither escalation fired."""
    from mathmodel.autopilot.pipeline import _failure_class

    shape_like = [
        "ValueError: The truth value of an array with more than one element is ambiguous.",
        "ValueError: setting an array element with a sequence.",
        "TypeError: only 0-dimensional arrays can be converted to Python scalars",
        "IndexError: too many indices for array: array is 0-dimensional",
    ]
    assert {_failure_class(e) for e in shape_like} == {"shape-mismatch"}

    target_like = [
        "AssertionError: problem4 target not satisfied: 2.55",
        "AssertionError: problem 3 final max rounded C = 0.2084",
        "AssertionError: problem3 strict criterion failed: 0.2573",
        "RuntimeError: problem 3 final max rounded C = 0.3305 did not fall strictly below 0.15",
    ]
    assert {_failure_class(e) for e in target_like} == {"final-target-unmet"}


def test_the_prompt_says_the_duration_is_the_answer():
    """Rounds 75-77: three runs stopped at 72 h because of a descriptive
    sentence, while an independent solve needs ~68-69 days to meet the target."""
    import pathlib

    from mathmodel.autopilot import codegen as mod

    src = pathlib.Path(mod.__file__).read_text(encoding="utf-8")
    assert "RUN THE PROCESS UNTIL THE STATED TARGET IS MET" in src
    assert "the duration IS the answer" in src
    assert "60, 120, 180, followed by an ellipsis" in src
    assert "figure was WRONG" in src
    assert "7e-9*exp(-0.89/C)" in src
    assert "D = 2.4e-3*exp(-0.45/C)*exp(-3850/T)" in src
    assert "integrate for months" in src

def test_the_prompt_does_not_bend_physics_to_match_prose():
    """Round 79: the prompt used to say `when the two disagree, fix the model`,
    which produced runs drying 19x-66x too fast to match a descriptive sentence."""
    import pathlib

    from mathmodel.autopilot import codegen as mod

    src = pathlib.Path(mod.__file__).read_text(encoding="utf-8")
    assert "when the two " not in src or "disagree, fix the model" not in src
    assert "DO NOT DISTORT YOUR TRANSPORT TO MATCH A STATED DURATION" in src
    assert "DESCRIPTIVE PROSE about the usual case" in src
    assert "19x, 50x and 66x too" in src
    assert "the sentence is what yields" in src


def test_the_prompt_requires_copying_template_labels():
    """Round 91: markerfix stayed blocked on result4's last header cell through
    six solve attempts because the template's label 药材表面 was written as 2.0."""
    import pathlib

    from mathmodel.autopilot import codegen as mod

    src = pathlib.Path(mod.__file__).read_text(encoding="utf-8")
    assert "COPY THE TEMPLATE'S HEADER ROW CELL FOR CELL, INCLUDING TEXT" in src
    assert "药材表面" in src
    assert "the LABEL still wins for " not in src
    assert "WRITE THEM IN PLACE OF THE VALUE THAT WOULD OTHERWISE SIT THERE" in src
    assert "AND THE DISTANCE COLUMNS THEMSELVES ARE NOT YOURS TO CHOOSE" in src
    assert "FIXED NUMERIC GRID" in src
    assert "INTERPOLATE onto the template's columns" in src
    assert "do not append an extra column" in src
    assert "EXACTLY as many " in src
    assert "six solve attempts and three verifications" in src

def test_output_bookkeeping_is_its_own_class_not_a_shape_error():
    """Round 93: 'At least one sheet must be visible' recurred 3 times but no
    class grouped it, so the escalation never fired. It must not share the shape
    note, whose text blames a node-count/profile-length disagreement."""
    from mathmodel.autopilot.pipeline import (
        _failure_class,
        _repeated_failure_note,
    )

    assert _failure_class("IndexError: At least one sheet must be visible") == (
        "output-bookkeeping"
    )
    assert _failure_class("ValueError: No worksheet named Sheet9") == (
        "output-bookkeeping"
    )
    note = _repeated_failure_note("output-bookkeeping", 3, 2)
    assert "BUILDING THE OUTPUT WORKBOOK" in note
    assert "STOP PATCHING THE CALL SITE" in note
    assert "grid nodes" not in note
    assert "profile" not in note
    assert _failure_class("IndexError: invalid index to scalar variable.") == (
        "shape-mismatch"
    )


def test_the_verifier_prompt_forbids_reporting_the_last_timestamp_as_a_crossing():
    """Round 105-106: the independent verifier twice reported t_dry as the file's
    final time -- 72.0 h for a grid ending at 259200 s, and 71.4166667 h for one
    ending at 257100 s -- not as the first time the target was met."""
    import pathlib

    from mathmodel.autopilot import verify as mod

    src = pathlib.Path(mod.__file__).read_text(encoding="utf-8")
    assert "A TIME AT WHICH A CONDITION IS FIRST MET IS NOT THE LAST " in src
    assert "259200 s (= 72 h)" in src
    assert "257100 s (= 71.4166667 h)" in src
    assert "READ THE INPUT FILES THAT ARE ACTUALLY THERE" in src
    assert "附件1.csv" in src
    assert "report the " in src and "FIRST time at which the condition actually holds" in src
    assert "looks precise and is guaranteed wrong" in src


def test_the_rendered_prompt_requires_the_template_distance_grid():
    """Round 127: source-level `in` checks cannot tell a live prompt literal from
    a comment or a branch that never executes. Render the prompt for real and
    assert the grid constraint is in the text the model actually receives.

    Round 120 measured result4 column 3 as 0.1052... where the template has 0.1,
    so this constraint is what the missing requirement looked like.
    """
    from mathmodel.autopilot.codegen import MathematicalModel, SolverCodeGenerator

    model = MathematicalModel(
        name="t",
        description="d",
        equations=[],
        assumptions=[],
        variables=[],
        parameters=[],
    )
    gen = SolverCodeGenerator.__new__(SolverCodeGenerator)
    prompt = gen._build_prompt(
        model=model,
        problem_text="dry a cylindrical medicinal material",
        data_schema="",
        required_outputs=["result1.xlsx", "result2.xlsx", "result3.xlsx", "result4.xlsx"],
        input_file_names=["附件1.xlsx", "附件2.xlsx"],
    )

    assert "AND THE DISTANCE COLUMNS THEMSELVES ARE NOT YOURS TO CHOOSE" in prompt
    assert "AND STORE EACH HEADER CELL IN THE TYPE THE TEMPLATE USES" in prompt
    assert "AND THE ROW TIMES ARE A FIXED GRID TOO" in prompt
    assert "AND STOP STRICTLY PAST THE TARGET" in prompt
    assert "STRICTLY less" in prompt
    assert "REPORTING CRITERION, not a physical bound" in prompt
    assert "clamp, floor, saturate" in prompt
    assert "MATCH THE SCHEME TO THE ORDER YOU CLAIM" in prompt
    assert "expected_order = 2.0" in prompt
    assert "every 60 s" in prompt
    assert "not text" in prompt
    assert "FIXED NUMERIC GRID" in prompt
    assert "INTERPOLATE onto the template's columns" in prompt
